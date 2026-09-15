"""Media and recall lifecycle against isolated SQLite/temp files only."""
import io
from zipfile import ZipFile

import pytest
from docx import Document
from sqlalchemy import event
from sqlalchemy.dialects import mysql

from tests.test_shipping_inspection import _user, _pc_client, _mini_client, _upload, _qr, storage, product_display_source
from app.shipping_inspection import service, file_service
from app.shipping_inspection.models import ShippingInspection

VIDEO = b"\x00\x00\x00\x18ftypisom\x00\x00\x00\x00isommp42" + b"test-video" * 4


def upload_video(client, item=None, version=0, content=VIDEO, filename="clip.mp4", mime="video/mp4"):
    data = {"outbound_record_id": "OB001", "edit_version": version}
    if item:
        data["item_id"] = item
    return client.post("/api/mini/shipping-inspection/videos", data=data, files={"file": (filename, content, mime)})


def submit(client, version=0, remark="保留备注"):
    return client.post("/api/mini/shipping-inspection/submit", json={"outbound_record_id": "OB001", "request_id": "request", "edit_version": version, "remark": remark})


def test_recall_retains_media_rejects_stale_requests_and_allows_resubmit(db, storage):
    user = _user(db)
    with _mini_client(db, user) as mini:
        photo = _upload(mini).json()
        video = upload_video(mini, "IT001").json()
        result = submit(mini)
        assert result.status_code == 200
        inspection_id = result.json()["id"]
        assert upload_video(mini).status_code == 400
        with _pc_client(db, user, ["shipping_inspection:read"]) as pc:
            assert pc.post(f"/api/shipping-inspection/records/{inspection_id}/recall", json={"edit_version": 0}).status_code == 403
        with _pc_client(db, user, ["shipping_inspection:write"]) as pc:
            for _ in range(2):
                recall = pc.post(f"/api/shipping-inspection/records/{inspection_id}/recall", json={"edit_version": 0})
                assert recall.status_code == 200
                assert recall.json()["data"]["edit_version"] == 1
            assert pc.get("/api/shipping-inspection/records").json()["data"]["total"] == 0
        scan = mini.post("/api/mini/shipping-inspection/scan", json={"qr_raw": _qr()}).json()
        assert scan["inspection"]["status"] == "draft"
        assert scan["inspection"]["remark"] == "保留备注"
        assert scan["inspection"]["edit_version"] == 1
        assert scan["photos"][0]["id"] == photo["id"]
        assert scan["videos"][0]["id"] == video["id"]
        assert (storage / video["file_path"]).is_file()
        assert submit(mini).status_code == 400
        assert upload_video(mini).status_code == 400
        assert mini.delete(f"/api/mini/shipping-inspection/photos/{photo['id']}").status_code == 400
        assert upload_video(mini, version=1).status_code == 200
        new_photo = mini.post(
            '/api/mini/shipping-inspection/photos',
            data={'outbound_record_id': 'OB001', 'edit_version': 1, 'item_id': 'IT001'},
            files={'file': ('after-recall.jpg', b'jpeg-bytes', 'image/jpeg')},
        )
        assert new_photo.status_code == 200
        assert submit(mini, 1, "").status_code == 200
        with _pc_client(db, user, ["shipping_inspection:write"]) as pc:
            assert pc.post(f"/api/shipping-inspection/records/{inspection_id}/recall", json={"edit_version": 0}).status_code == 409
            detail = pc.get(f"/api/shipping-inspection/records/{inspection_id}").json()["data"]
            assert len(detail["photos"]) == 2
            assert len(detail["videos"]) == 2
            assert detail["remark"] == ""
    inspection = db.get(ShippingInspection, inspection_id)
    assert inspection.recalled_by == user.id
    assert inspection.recalled_at is not None
    assert inspection.photo_count == 2


def test_failed_recall_commit_rolls_back_state(db, storage, monkeypatch):
    user = _user(db)
    with _mini_client(db, user) as mini:
        _upload(mini)
        inspection_id = submit(mini).json()['id']
    def fail_commit():
        raise RuntimeError('simulated database commit failure')
    monkeypatch.setattr(db, 'commit', fail_commit)
    with pytest.raises(RuntimeError, match='simulated database'):
        service.recall(db, inspection_id, user.id, 0)
    inspection = db.get(ShippingInspection, inspection_id)
    assert inspection.status == 'submitted'
    assert inspection.edit_version == 0
    assert inspection.recalled_at is None


def test_video_is_private_separate_from_photos_and_never_satisfies_submit(db, storage):
    user = _user(db)
    with _mini_client(db, user) as mini:
        video = upload_video(mini).json()
        scan = mini.post("/api/mini/shipping-inspection/scan", json={"qr_raw": _qr()}).json()
        assert scan["photos"] == []
        assert scan["inspection"]["photo_count"] == 0
        assert scan["videos"][0]["media_type"] == "video"
        assert submit(mini).status_code == 400
        response = mini.get(f"/api/mini/shipping-inspection/images/{video['file_path']}")
        assert response.content == VIDEO
        assert response.headers['content-type'].startswith('video/mp4')
        assert mini.delete(f"/api/mini/shipping-inspection/photos/{video['id']}").status_code == 400
        assert mini.delete(f"/api/mini/shipping-inspection/videos/{video['id']}").status_code == 200
        assert not (storage / video['file_path']).exists()
        # HTTPBearer blocks a missing credential with 403 before route execution.
        assert mini.get('/api/mini/shipping-inspection/images/xx/missing.mp4', headers={'Authorization': ''}).status_code == 403
    with _pc_client(db, user, [], roles=['super_admin']) as pc:
        rows = pc.get('/api/shipping-inspection/outbound-records').json()['data']['items']
        assert next(row for row in rows if row['outbound_record_id'] == 'OB001')['photo_count'] == 0


@pytest.mark.parametrize('filename,mime,content', [('clip.exe','video/mp4',VIDEO), ('clip.mp4','image/png',VIDEO), ('clip.mp4','video/mp4',b'not-video')])
def test_invalid_video_leaves_no_file(db, storage, filename, mime, content):
    with _mini_client(db, _user(db)) as client:
        assert upload_video(client, filename=filename, mime=mime, content=content).status_code == 400
    assert not list(storage.rglob('*.*'))


def test_video_size_limit_and_unknown_item_do_not_leave_files(db, storage, monkeypatch):
    with _mini_client(db, _user(db)) as client:
        assert upload_video(client, item='IT003').status_code == 400
        monkeypatch.setattr(file_service, 'VIDEO_MAX_BYTES', 30)
        assert upload_video(client).status_code == 400
    assert not list(storage.rglob('*.*'))


def test_photo_count_uses_current_read_after_the_inspection_lock(db, storage):
    user = _user(db)
    with _mini_client(db, user) as client:
        _upload(client)
    statements = []
    def capture(state):
        if state.is_select:
            statements.append(str(state.statement.compile(dialect=mysql.dialect())))
    event.listen(db, 'do_orm_execute', capture)
    try:
        service.submit(db, outbound_record_id='OB001', user_id=user.id)
    finally:
        event.remove(db, 'do_orm_execute', capture)
    reads = [sql for sql in statements if 'FROM ark_shipping_inspection_photos' in sql]
    assert reads and all('FOR UPDATE' in sql for sql in reads)


def test_word_download_preserves_layout_qr_and_scope(db):
    user = _user(db)
    with _pc_client(db, user, ['shipping_inspection:read']) as client:
        assert client.get('/api/shipping-inspection/outbound-records/OB001/word').status_code == 422
    with _pc_client(db, user, [], roles=['super_admin']) as client:
        response = client.get('/api/shipping-inspection/outbound-records/OB001/word')
    assert response.status_code == 200
    assert response.headers['content-type'].endswith('wordprocessingml.document')
    assert '.docx' in response.headers['content-disposition']
    document = Document(io.BytesIO(response.content))
    section = document.sections[0]
    assert abs(section.page_width.mm - 210) < 0.1
    assert abs(section.left_margin.mm - 6) < 0.1
    table = document.tables[-1]
    assert [cell.text for cell in table.rows[0].cells] == ['#','产品类别','规格','颜色/尺寸/克重','数量','批次号']
    assert table.cell(1, 5).text == ''
    assert len(document.inline_shapes) == 1
    with ZipFile(io.BytesIO(response.content)) as package:
        xml = package.read('word/document.xml').decode()
    assert 'F5F5F5' in xml and 'w:val="center"' in xml and 'w:tblLayout w:type="fixed"' in xml
