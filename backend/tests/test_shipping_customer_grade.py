import io
from datetime import date
from decimal import Decimal

import pytest

from docx import Document
from sqlalchemy import text

from app.invoice.customer_profile_service import save_customer_grade
from app.invoice.models import Invoice
from app.shipping_inspection import outbound_service
from app.shipping_inspection.print_customer_service import with_customer_order_info
from app.shipping_inspection.word_service import build_outbound_word


def seed(db):
    outbound_service._columns_cache.clear()
    db.execute(text("ALTER TABLE lsordertest.okki_outbound_record_items ADD COLUMN order_id TEXT"))
    db.execute(text("UPDATE lsordertest.okki_outbound_record_items SET order_id='O1' WHERE id IN ('IT001','IT002')"))
    db.execute(text("INSERT INTO lsordertest.okki_orders(order_id, company_id, amount_usd) VALUES ('O1','C1',123.45)"))
    save_customer_grade(db, "C1", "S", None)
    return {"company_id": "C1", "outbound_record_id": "OB001", "customer_name": "Customer"}


def test_exact_order_deduplicated_and_word_matches(db):
    record = seed(db)
    result = with_customer_order_info(db, record)
    assert result["customer_grade"] == "S"
    assert result["order_amount_text"] == "USD 123.45"
    doc = Document(io.BytesIO(build_outbound_word(result, [], "test")))
    assert doc.tables[0].cell(0, 2).text == "客户等级\nS"
    assert doc.tables[0].cell(0, 3).text == "订单金额\nUSD 123.45"
    assert "customer_grade" not in record


def test_local_order_keeps_currency_zero_and_customer_boundary(db):
    record = seed(db)
    db.add(Invoice(invoice_no="I1", customer_id="C1", customer_name="Customer",
                   invoice_date=date(2026,9,18), currency="EUR", sync_status="synced", total_amount=Decimal("0"), xiaoman_order_id="O1"))
    db.flush()
    assert with_customer_order_info(db, record)["order_amount_text"] == "EUR 0.00"
    assert with_customer_order_info(db, {**record, "company_id": "unrelated"})["order_amount_text"] == "—"


def test_missing_order_never_prints_partial_sum(db):
    record = seed(db)
    db.execute(text("UPDATE lsordertest.okki_outbound_record_items SET order_id='missing' WHERE id='IT002'"))
    assert with_customer_order_info(db, record)["order_amount_text"] == "—"


def test_multiple_currencies_not_added_together(db):
    record = seed(db)
    db.execute(text("UPDATE lsordertest.okki_outbound_record_items SET order_id='O2' WHERE id='IT002'"))
    db.add(Invoice(invoice_no="I2", customer_id="C1", customer_name="Customer",
                   invoice_date=date(2026,9,18), currency="EUR", sync_status="synced", total_amount=Decimal("50"), xiaoman_order_id="O2"))
    db.flush()
    assert with_customer_order_info(db, record)["order_amount_text"] == "EUR 50.00 / USD 123.45"


@pytest.mark.parametrize("order_id", [None, "", "   "])
def test_null_order_link_never_prints_partial_sum(db, order_id):
    record = seed(db)
    db.execute(text("UPDATE lsordertest.okki_outbound_record_items SET order_id=:order_id WHERE id='IT002'"), {"order_id": order_id})
    assert with_customer_order_info(db, record)["order_amount_text"] == "—"


@pytest.mark.parametrize("sync_status", ["not_synced", "sync_failed", "sync_uncertain"])
def test_unsynced_local_edit_does_not_replace_confirmed_order_amount(db, sync_status):
    record = seed(db)
    db.add(Invoice(invoice_no="I1", customer_id="C1", customer_name="Customer",
                   invoice_date=date(2026,9,18), currency="EUR", sync_status=sync_status,
                   total_amount=Decimal("999"), xiaoman_order_id="O1"))
    db.flush()
    assert with_customer_order_info(db, record)["order_amount_text"] == "USD 123.45"


def test_print_api_and_word_use_invoice_bridge_and_same_metadata(db, monkeypatch):
    from app.shipping_inspection import router
    from tests.test_shipping_inspection import _pc_client, _user

    seed(db)
    db.execute(text("CREATE TABLE lsordertest.okki_products (product_id INTEGER PRIMARY KEY, model TEXT, size TEXT, color TEXT)"))
    db.execute(text("ALTER TABLE lsordertest.okki_outbound_records ADD COLUMN company_id TEXT"))
    db.execute(text("ALTER TABLE lsordertest.okki_outbound_records ADD COLUMN outbound_invoice_id TEXT"))
    db.execute(text("ALTER TABLE lsordertest.okki_outbound_record_items ADD COLUMN outbound_invoice_id TEXT"))
    db.execute(text("UPDATE lsordertest.okki_outbound_records SET company_id='C1', outbound_invoice_id='BRIDGE' WHERE id='OB001'"))
    db.execute(text("UPDATE lsordertest.okki_outbound_record_items SET outbound_invoice_id='BRIDGE', outbound_record_id='OTHER' WHERE id IN ('IT001','IT002')"))
    outbound_service._columns_cache.clear()
    monkeypatch.setattr(router, "with_owner_chinese_name", lambda db, record: record)
    with _pc_client(db, _user(db), [], roles=["super_admin"]) as client:
        response = client.get("/api/shipping-inspection/outbound-records/OB001/print-data")
        assert response.status_code == 200
        record = response.json()["data"]["record"]
        assert record["customer_grade"] == "S"
        assert record["order_amount_text"] == "USD 123.45"
        response = client.get("/api/shipping-inspection/outbound-records/OB001/word")
        assert response.status_code == 200
        doc = Document(io.BytesIO(response.content))
        assert doc.tables[0].cell(0, 3).text == "订单金额\nUSD 123.45"
    outbound_service._columns_cache.clear()
