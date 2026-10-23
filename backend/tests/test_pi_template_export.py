"""Employee PI template exports: source values, artwork and pagination."""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile
import re

from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from app.invoice import pi_template


def invoice(count=1, **changes):
    line = dict(product_display="Standard Double Drawn Genius Weft", product_name="Hair full name",
        model="", curl="", color="#2TP2/6", length="16", net_weight_grams="20g", quantity=3,
        price_per_piece=Decimal("19.5300"), discount_amount=Decimal("-1.23"), total_price=Decimal("57.36"),
        product_kind="hair")
    data = dict(order_type="stock", invoice_no="PI-TEMPLATE-001", invoice_date=date(2026, 10, 23),
        customer_name="Template Buyer", contact_name="Recipient", contact_phone="+1 555 123",
        contact_email="buyer@example.test", delivery_address="25 Test Street\nSample City",
        sales_user_name="Alice", sales_phone="+86 123", sales_email="alice@example.test", currency="EUR",
        items=[SimpleNamespace(**deepcopy(line)) for _ in range(count)], product_amount=Decimal("57.36") * count,
        internal_accessory=Decimal("2.10"), packaging_quantity=2, shipping_fee=Decimal("40.00"),
        surcharge_name="Payment Surcharge", surcharge_amount=Decimal("7.51"), payment_term="100% prepaid",
        express_channel="DHL", remark="Please confirm delivery.")
    data["total_amount"] = data["product_amount"] + Decimal("49.61")
    return SimpleNamespace(**{**data, **changes})


@pytest.mark.parametrize("count", [0, 1, 4, 5, 6, 65])
def test_excel_uses_template_and_exact_stored_values(count):
    source = invoice(count)
    result = pi_template.build_workbook(source)
    wb = load_workbook(result, data_only=False)
    ws = wb.active
    assert len(wb.worksheets) == 1
    assert ws["C6"].value == "Template Buyer / Recipient"
    assert ws["G6"].value == "Alice"
    assert ws["C7"].value == "'+1 555 123"
    assert ws["G8"].hyperlink is None
    assert ws["G9"].value.date() == source.invoice_date
    assert ws["G10"].value == source.invoice_no
    assert ws["B11"].value == "Product"
    assert ws["F12"].value == (3 if count else None)
    assert ws["G12"].value == (19.53 if count else None)
    assert ws["H12"].value == (57.36 if count else None)
    if count:
        assert "Discount: -1.23 EUR" in ws["B12"].value
        assert ws.cell(11 + count, 1).value == count
    cells = {str(c.value): ws.cell(c.row, 8).value for row in ws for c in row if c.column == 2 and c.value}
    assert cells["Hair Cost"] == float(source.product_amount)
    assert cells["Shipping Cost"] == 40
    assert cells["Payment Surcharge"] == 7.51  # Never the sample's hardcoded 5%.
    assert cells["Packaging (2)"] == 2.1
    assert cells["Total (EUR)"] == float(source.total_amount)
    assert "C9:E10" in str(ws.merged_cells)
    assert ws.print_title_rows == "$1:$11"
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert not any(c.data_type == "f" for row in ws for c in row)
    with ZipFile(result) as out, ZipFile(pi_template.TEMPLATE_PATH) as original:
        assert out.read("xl/drawings/drawing1.xml") == original.read("xl/drawings/drawing1.xml")
        assert out.read("xl/media/image1.png") == original.read("xl/media/image1.png")
        assert b"derek@" not in b"".join(out.read(n).lower() for n in out.namelist() if n.endswith(".xml"))


def test_no_sample_identity_formula_or_internal_financial_data():
    source = invoice(customer_name="=HYPERLINK(\"evil\")", contact_name="", sales_user_name=None,
        sales_phone=None, sales_email=None, contact_email=None, payment_term=None, remark="@unsafe",
        internal_received="INTERNAL RECEIVED SECRET", internal_balance="INTERNAL BALANCE SECRET")
    wb = load_workbook(pi_template.build_workbook(source))
    ws = wb.active
    assert ws["C6"].value.startswith("'=HYPERLINK")
    assert all(ws[address].value is None for address in ["G6", "G7", "G8", "C8"])
    values = " ".join(str(c.value) for row in ws for c in row if c.value)
    assert "SECRET" not in values and "PayPal" not in values and "Derek" not in values
    assert "'@unsafe" in values


def test_mixed_accessory_and_high_precision_transaction_price():
    source = invoice()
    source.items[0].product_kind = "accessory"
    source.items[0].price_per_piece = Decimal("19.5301")
    ws = load_workbook(pi_template.build_workbook(source)).active
    assert ws["G12"].value == 19.5301
    assert ws["G12"].number_format == "0.00##"
    assert any(c.value == "Products Cost" for row in ws for c in row)


@pytest.mark.parametrize("count", [1, 65])
def test_pdf_has_template_header_real_values_and_all_lines(count):
    source = invoice(count)
    for i, line in enumerate(source.items):
        line.product_display = f"Product-{i + 1:03d}"
    pdf = pi_template.build_pdf(source).getvalue()
    reader = PdfReader(BytesIO(pdf))
    content = "\n".join(page.extract_text() for page in reader.pages)
    assert "JUANCHENG LESHINE" in content
    assert "Proforma" in content and "Invoice" in content
    assert "Template Buyer" in content and "Alice" in content
    assert "57.36" in content and "7.51" in content and "100% prepaid" in content
    assert "Total (EUR)" in content
    assert f"Product-{count:03d}" in content
    assert "Derek" not in content
    assert len(reader.pages) > 1 if count > 50 else len(reader.pages) == 1


def test_long_multilingual_fields_wrap_without_truncating():
    source = invoice(delivery_address="上海市完整收货地址 " * 30 + "ADDRESS-END", remark="备注 " * 100 + "NOTE-END")
    source.items[0].product_display = "长商品名称 " * 25 + "PRODUCT-END"
    source.items[0].color = "#LONG-COLOR-" * 12 + "COLOR-END"
    pdf = PdfReader(pi_template.build_pdf(source))
    text = "\n".join(p.extract_text() for p in pdf.pages)
    # A narrow color cell can wrap inside the marker; every character must remain.
    normalized = re.sub(r"\s+", "", text)
    assert all(marker in normalized for marker in ["ADDRESS-END", "NOTE-END", "PRODUCT-END", "COLOR-END"])
