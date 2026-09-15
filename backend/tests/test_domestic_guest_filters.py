import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from alembic.migration import MigrationContext
from alembic.operations import Operations
from sqlalchemy import create_engine, inspect, text

from app.domestic import customer_service, order_service
from app.domestic.models import DomesticOrder
from app.domestic.schemas import OrderCreate, OrderUpdate, OrderItemUpdate
from app.domestic.export_service import build_order_workbook
from tests.test_domestic_customer_order_controls import _user, _customer
from tests.test_domestic_order_channel_source import context
from tests.test_domestic_export import _order_detail


def test_combined_customer_filters_and_pagination(db):
    owner, other = _user(db, "owner"), _user(db, "other")
    first = _customer(db, owner, "first")
    second = _customer(db, owner, "second")
    third = _customer(db, other, "third")
    first.customer_level, second.customer_level, third.customer_level = "A", "B", "A"
    db.flush()
    rows, total = customer_service.list_customers(db, customer_level="A", owner_user_id=owner.id, province="山东省", owner_scope="private")
    assert total == 1 and rows[0]["id"] == first.id
    assert customer_service.list_customers(db, customer_level="A", page_size=1)[1] == 2
    assert customer_service.list_customers(db, owner_user_id=owner.id)[1] == 2
    assert customer_service.list_customers(db, owner_user_id=owner.id, owner_scope="public")[1] == 0
    assert customer_service.list_customers(db)[1] == 3


def test_guest_create_detail_edit_clear_and_replay(db):
    user, customer, payload = context(db, "guest-field")
    data = payload.model_dump()
    data["items"][0]["guest_name"] = "  王女士  "
    request = OrderCreate.model_validate(data)
    result = order_service.create_order(db, request, user.id)
    order_id = result["id"]
    detail = order_service.get_order_detail(db, order_id)
    item_id = detail["items"][0]["id"]
    assert "guest_name" not in detail
    assert detail["items"][0]["guest_name"] == "王女士"
    assert order_service.create_order(db, request, user.id)["replayed"]
    order_service.update_item(db, item_id, OrderItemUpdate(guest_name=" 李先生 "), user.id)
    assert order_service.get_order_detail(db, order_id)["items"][0]["guest_name"] == "李先生"
    order_service.update_item(db, item_id, OrderItemUpdate(guest_name="  "), user.id)
    assert order_service.get_order_detail(db, order_id)["items"][0]["guest_name"] is None
    with pytest.raises(ValidationError):
        OrderItemUpdate(guest_name="名" * 121)
    with pytest.raises(ValidationError):
        OrderUpdate(guest_name="不再录入订单头")


def test_guest_printed_on_item_rows_and_absent_from_production():
    detail = _order_detail()
    detail["items"][0]["guest_name"] = "=王女士"
    wb = load_workbook(build_order_workbook(detail, "销售"))
    for ws in wb.worksheets:
        assert "顾客：" not in ws["A3"].value
        assert any("顾客：'=王女士" in str(cell.value) and cell.data_type == "s" for row in ws for cell in row)
    detail["order_kind"] = "production"
    wb = load_workbook(build_order_workbook(detail, "销售"))
    assert all("顾客：'=王女士" not in str(cell.value) for ws in wb for row in ws for cell in row)


def test_guest_migration_preserves_existing_rows():
    engine = create_engine("sqlite:///:memory:")
    path = Path(__file__).parents[1] / "alembic/versions/145_domestic_order_guest.py"
    spec = importlib.util.spec_from_file_location("guest_migration", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE ark_domestic_orders (id INTEGER PRIMARY KEY, order_no VARCHAR(64))"))
        conn.execute(text("INSERT INTO ark_domestic_orders VALUES (1, 'legacy')"))
        module.op = Operations(MigrationContext.configure(conn))
        module.upgrade()
        assert conn.execute(text("SELECT order_no, guest_name FROM ark_domestic_orders")).one() == ("legacy", None)
        assert next(c for c in inspect(conn).get_columns("ark_domestic_orders") if c["name"] == "guest_name")["nullable"]


@pytest.mark.parametrize("kind", ["business", "production"])
def test_empty_guest_preserves_existing_request_hash(db, kind):
    import hashlib
    import json
    user, customer, payload = context(db, "old-hash")
    data = payload.model_dump()
    data["order_kind"] = kind
    payload = OrderCreate.model_validate(data)
    old_data = payload.model_dump(mode="json", exclude={"request_id", "guest_name"})
    for item in old_data["items"]:
        item.pop("guest_name", None)
        item.pop("guest_order_date", None)
    old_hash = hashlib.sha256(json.dumps(old_data, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
    assert order_service._order_request_hash(payload) == old_hash
    if kind == "business":
        result = order_service.create_order(db, payload, user.id)
        db.get(DomesticOrder, result["id"]).request_hash = old_hash
        db.commit()
        assert order_service.create_order(db, payload, user.id)["replayed"]
        changed = payload.model_copy(deep=True)
        changed.items[0].guest_name = "New guest"
        assert order_service._order_request_hash(changed) != old_hash


def test_item_guest_migration_backfills_and_preserves_existing_names():
    engine = create_engine("sqlite:///:memory:")
    path = Path(__file__).parents[1] / "alembic/versions/150_domestic_item_guest.py"
    spec = importlib.util.spec_from_file_location("item_guest_migration", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE ark_domestic_orders (id INTEGER PRIMARY KEY, order_kind VARCHAR(20), guest_name VARCHAR(120))"))
        conn.execute(text("CREATE TABLE ark_domestic_order_items (id INTEGER PRIMARY KEY, order_id INTEGER)"))
        conn.execute(text("INSERT INTO ark_domestic_orders VALUES (1, 'business', '王女士'), (2, 'production', NULL), (3, 'business', NULL)"))
        conn.execute(text("INSERT INTO ark_domestic_order_items VALUES (1, 1), (2, 1), (3, 2), (4, 3)"))
        module.op = Operations(MigrationContext.configure(conn))
        module.upgrade()
        assert conn.execute(text("SELECT guest_order_date FROM ark_domestic_order_items")).scalars().all() == [None] * 4
        assert conn.execute(text("SELECT guest_name FROM ark_domestic_order_items ORDER BY id")).scalars().all() == ["王女士", "王女士", None, None]
        conn.execute(text("UPDATE ark_domestic_order_items SET guest_name='李先生' WHERE id=2"))
        module.upgrade()
        assert conn.execute(text("SELECT guest_name FROM ark_domestic_order_items WHERE id=2")).scalar() == "李先生"
        assert conn.execute(text("SELECT guest_name FROM ark_domestic_orders WHERE id=1")).scalar() == "王女士"


def test_public_progress_contract_excludes_quantities_and_money():
    detail = {"items": [{"id": 1, "guest_name": "王女士", "order_qty": 2,
        "unit_price": 900, "route_id": 3, "remark": "备注",
        "steps": [{"process_name": "植发", "passed_qty": 2}, {"process_name": "造型", "passed_qty": 1}]}]}
    item = order_service.track_public_view(detail)["items"][0]
    assert item["steps"] == [{"process_name": "植发", "completed": True}, {"process_name": "造型", "completed": False}]
    assert not {"order_qty", "unit_price", "route_id"} & item.keys()
    assert item["remark"] == "备注"


def test_empty_item_guest_preserves_append_replay(db):
    import hashlib
    import json
    from app.domestic.models import DomesticItemAppendRequest
    from app.domestic.schemas import OrderItemAppend
    user, customer, payload = context(db, "append-guest")
    payload.is_draft = True
    order = order_service.create_order(db, payload, user.id)
    append = OrderItemAppend.model_validate({**payload.items[0].model_dump(), "request_id": "append-guest-replay"})
    old_data = append.model_dump(mode="json", exclude={"request_id", "guest_name", "guest_order_date"})
    old_hash = hashlib.sha256(json.dumps(old_data, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
    assert order_service._item_append_request_hash(append) == old_hash
    first = order_service.add_item(db, order["id"], append, user.id)
    row = db.query(DomesticItemAppendRequest).filter_by(request_id=append.request_id).one()
    row.request_hash = old_hash
    db.commit()
    assert order_service.add_item(db, order["id"], append, user.id)["id"] == first["id"]
    changed = append.model_copy(update={"guest_name": "新顾客"})
    assert order_service._item_append_request_hash(changed) != old_hash


@pytest.mark.parametrize("raw, expected", [
    ("2026-09-15", "2026-09-15"),
    ("2026-09-14", "2026-09-14"),
    ("2028-02-29", "2028-02-29"),
])
def test_guest_order_date_create_update_and_public_view(db, raw, expected):
    from app.domestic.models import DomesticOrderItem
    user, customer, payload = context(db, "guest-order-time")
    data = payload.model_dump()
    data["items"][0]["guest_order_date"] = raw
    request = OrderCreate.model_validate(data)
    result = order_service.create_order(db, request, user.id)
    detail = order_service.get_order_detail(db, result["id"])
    item = detail["items"][0]
    assert item["guest_order_date"] == expected
    from datetime import date
    assert type(db.get(DomesticOrderItem, item["id"]).guest_order_date) is date
    assert order_service.track_public_view(detail)["items"][0]["guest_order_date"] == expected
    order_service.update_item(db, item["id"], OrderItemUpdate(guest_order_date="2026-09-16"), user.id)
    assert order_service.get_order_detail(db, result["id"])["items"][0]["guest_order_date"] == "2026-09-16"
    order_service.update_item(db, item["id"], OrderItemUpdate(guest_order_date=None), user.id)
    assert db.get(DomesticOrderItem, item["id"]).guest_order_date is None


def test_guest_order_date_rejects_invalid_input():
    with pytest.raises(ValidationError):
        OrderItemUpdate(guest_order_date="不是日期")


@pytest.mark.parametrize("kind", ["business", "production"])
def test_guest_order_date_append_and_production_exclusion(db, kind):
    from app.domestic.models import DomesticOrderItem
    from app.domestic.schemas import OrderItemAppend
    user, customer, payload = context(db, "time-append")
    data = payload.model_dump()
    data.update(order_kind=kind, is_draft=True)
    data["items"][0]["guest_order_date"] = "2026-09-15"
    payload = OrderCreate.model_validate(data)
    order = order_service.create_order(db, payload, user.id)
    first = order_service.get_order_detail(db, order["id"])["items"][0]
    assert first["guest_order_date"] == ("2026-09-15" if kind == "business" else None)
    append = OrderItemAppend.model_validate({**payload.items[0].model_dump(),
        "request_id": "time-append-request", "guest_order_date": "2026-09-14"})
    result = order_service.add_item(db, order["id"], append, user.id)
    saved = db.get(DomesticOrderItem, result["id"])
    if kind == "business":
        assert saved.guest_order_date.isoformat() == "2026-09-14"
    else:
        assert saved.guest_order_date is None
        with pytest.raises(ValueError):
            order_service.update_item(db, saved.id, OrderItemUpdate(guest_order_date="2026-09-15"), user.id)
