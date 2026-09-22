"""Print-only ordering and selective DOCX emphasis; no production database."""
import io

from docx import Document

from app.shipping_inspection.print_service import (
    annotate_print_items,
    is_other_accessory,
    sort_outbound_print_items,
)
from app.shipping_inspection.word_service import build_outbound_word
from tests.test_shipping_inspection import _user, _pc_client


def sample_items():
    return [
        {"item_id": "b3p", "spec": "B3平行", "size": "16"},
        {"item_id": "24", "spec": "B1天才", "size": "24"},
        {"item_id": "unknown", "spec": "B1天才", "size": "未知"},
        {"item_id": "b1p", "spec": "B1平行", "size": "16"},
        {"item_id": "22", "spec": "B1天才", "product_name": "Weft／22／#8/20g"},
        {"item_id": "18", "spec": "B1天才", "size": '18"'},
        {"item_id": "b3t", "spec": "B3天才", "size": "24"},
        {"item_id": "20", "spec": "B1天才", "size": "20寸"},
        {"item_id": "16", "spec": "B1天才", "size": "16"},
        {"item_id": "16b", "spec": "B1天才", "size": "16"},
    ]


EXPECTED = ["16", "16b", "18", "20", "22", "24", "unknown", "b1p", "b3t", "b3p"]


def test_sort_groups_then_numeric_sizes_and_preserves_source():
    items = sample_items()
    original = list(items)
    assert [i["item_id"] for i in sort_outbound_print_items(items)] == EXPECTED
    assert items == original
    assert [i["size"] for i in sort_outbound_print_items([
        {"size": "100"}, {"size": "9.5"}, {"size": "9"}, {"size": None},
    ])] == ["9", "9.5", "100", None]
    assert [i.get("spec") for i in sort_outbound_print_items([
        {}, {"spec": "B10"}, {"spec": "B3"}, {"spec": "B1"},
    ])] == ["B1", "B3", "B10", None]


def test_word_order_and_only_grade_is_large_bold():
    items = sample_items()
    for item in items:
        item["product_name"] = item.get("product_name", item["item_id"])
    document = Document(io.BytesIO(build_outbound_word({}, items, "test")))
    rows = document.tables[-1].rows[1:-1]  # 去掉表头与合计行
    assert [r.cells[1].text for r in rows] == ["16", "16b", "18", "20", "Weft", "24", "unknown", "b1p", "b3t", "b3p"]
    assert [r.cells[0].text for r in rows] == [str(i) for i in range(1, 11)]
    for row in rows:
        runs = [run for run in row.cells[2].paragraphs[0].runs if run.text]
        assert runs[0].text in ("B1", "B3")
        assert runs[0].bold and runs[0].font.size.pt == 12
        assert not runs[1].bold and runs[1].font.size.pt == 9.5  # DOCX stores half-point units, matching the existing body style.
    assert document.tables[-1].rows[-1].cells[0].text == "合计"
    assert document.tables[-1].rows[-1].cells[4].text == "0"  # sample_items 未填 qty
    document = Document(io.BytesIO(build_outbound_word({}, [{"spec": "B10 AB1 B3平行"}], "test")))
    runs = document.tables[-1].cell(1, 2).paragraphs[0].runs
    assert [r.text for r in runs if r.bold] == ["B3"]


def test_print_api_and_word_share_order(db, monkeypatch):
    from app.shipping_inspection import outbound_service
    monkeypatch.setattr(outbound_service, "list_outbound_items", lambda *args: sample_items())
    with _pc_client(db, _user(db), [], roles=["super_admin"]) as client:
        response = client.get("/api/shipping-inspection/outbound-records/OB001/print-data")
        assert response.status_code == 200
        assert [i["item_id"] for i in response.json()["data"]["items"]] == EXPECTED


def test_word_masks_customer_name_without_changing_source():
    for name, expected in [
        ("Inessa Wassiljev/Haarverlängerung", "Ine***"), ("AB", "AB***"),
        ("ABC", "ABC***"), ("王女士旗舰店", "王女士***"),
        ("😀AB Customer", "😀AB***"), (" <&>Company ", "<&>***"),
        (None, ""), ("", ""), ("   ", ""),
    ]:
        record = {"customer_name": name}
        document = Document(io.BytesIO(build_outbound_word(record, [], "test")))
        assert document.tables[0].cell(0, 1).text == expected
        assert record["customer_name"] == name


def test_is_other_accessory_only_matches_accessory_named_other():
    assert is_other_accessory({"product_kind": "accessory", "product_name": "Other"})
    assert is_other_accessory({"product_kind": "accessory", "product_name": " other "})
    assert not is_other_accessory({"product_kind": "accessory", "product_name": "Hair Gripper"})
    assert not is_other_accessory({"product_kind": "hair", "product_name": "Other"})
    assert not is_other_accessory({"product_name": "Other"})


def test_annotate_print_items_marks_kind_and_drops_other_accessories(db):
    from app.invoice.models import Invoice, InvoiceItem, StdPrice
    db.query(StdPrice).delete()
    db.query(InvoiceItem).delete()
    db.query(Invoice).delete()
    db.add(StdPrice(product_kind="accessory", product_id=901, sku_id=1, accessory_name="Hair Gripper",
                    accessory_model="Tape", accessory_color="Black", currency="USD", price=0))
    db.add(StdPrice(product_kind="accessory", product_id=902, sku_id=2, accessory_name="Other",
                    accessory_model="Misc", accessory_color="—", currency="USD", price=0))
    from datetime import date
    invoice = Invoice(invoice_no="IV-PRINT-1", customer_id="C1", customer_name="C",
                      currency="USD", order_type="stock", product_amount=0, total_amount=0,
                      invoice_date=date(2026, 9, 1))
    db.add(invoice)
    db.flush()
    db.add(InvoiceItem(invoice_id=invoice.id, product_kind="accessory", product_id=903, product_name="Clip",
                       product_display="Clip", color="Black", quantity=1, total_price=0))
    db.commit()
    items = [
        {"item_id": "h", "product_id": 100, "product_name": "Weft/22", "qty": 2},
        {"item_id": "a", "product_id": 901, "product_name": "Hair Gripper", "qty": 5},
        {"item_id": "o", "product_id": 902, "product_name": "Other", "qty": 9},
        {"item_id": "i", "product_id": 903, "product_name": "Clip", "qty": 1},
    ]
    result = annotate_print_items(db, items)
    assert [i["item_id"] for i in result] == ["h", "a", "i"]
    assert [i["product_kind"] for i in result] == ["hair", "accessory", "accessory"]
    # 源明细不被就地修改
    assert "product_kind" not in items[0] and items[2]["item_id"] == "o"


def test_word_splits_product_accessory_tables_with_qty_summary():
    items = [
        {"product_name": "Weft/22/#1/20g", "spec": "B1天才", "qty": 2, "product_kind": "hair"},
        {"product_name": "Hair Gripper", "spec": "魔术贴", "qty": 3, "product_kind": "accessory"},
    ]
    document = Document(io.BytesIO(build_outbound_word({}, items, "test")))
    tables = document.tables
    assert len(tables) >= 3  # 头信息 + 产品表 + 配件表
    product_table, accessory_table = tables[-2], tables[-1]
    assert product_table.cell(0, 1).text == "产品类别"
    assert product_table.rows[-1].cells[0].text == "合计"
    assert product_table.rows[-1].cells[4].text == "2"
    assert accessory_table.rows[-1].cells[0].text == "合计"
    assert accessory_table.rows[-1].cells[4].text == "3"
    # 仅产品时不出现配件表
    only = Document(io.BytesIO(build_outbound_word({}, [items[0]], "test")))
    assert only.tables[-1].rows[-1].cells[0].text == "合计"
    assert all("配件明细" not in p.text for p in only.paragraphs)
