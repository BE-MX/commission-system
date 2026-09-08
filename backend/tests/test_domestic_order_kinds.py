"""Production ordering never depends on a customer/price and stops at storage."""

from datetime import date
from decimal import Decimal
import importlib.util
from pathlib import Path
from uuid import uuid4

import pytest
from pydantic import ValidationError

from app.auth.models import ArkUser
from app.domestic import constants as C, order_service, pricing_service, report_service, route_rule_service
from app.domestic.models import (
    DomesticCustomer, DomesticCustomerLedger, DomesticItemProgress, DomesticOrder,
    DomesticOrderItem, DomesticProduct, DomesticRouteRule,
)
from app.domestic.schemas import (
    DraftSubmitRequest, ItemShipRequest, OrderCreate, OrderItemAppend, OrderItemUpdate, OrderUpdate,
)
from app.production.models import Process, ProcessRoute, ProcessRouteStep, UserProcessBinding
from app.system.models import SysDict


def migration():
    path = Path(__file__).parents[1] / "alembic/versions/140_domestic_order_kinds.py"
    spec = importlib.util.spec_from_file_location("domestic_kinds_migration", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def context(db):
    user = ArkUser(username="production-order-user", password_hash="test", real_name="备货员")
    db.add(user)
    db.flush()
    process_names = ["确认下单", "毛坯质检", "毛坯维修", "入库", "毛坯出库", "做发型", "发货完成"]
    processes = []
    for name in process_names:
        process = Process(name=name)
        db.add(process)
        processes.append(process)
    db.flush()
    for name in C.DEFAULT_ROUTE_NAMES.values():
        route = ProcessRoute(name=name)
        db.add(route)
        db.flush()
        for index, process in enumerate(processes, 1):
            db.add(ProcessRouteStep(route_id=route.id, process_id=process.id, step_order=index))
        db.flush()
        db.add(DomesticRouteRule(route_id=route.id, process_id=processes[1].id, rule_type="decision", config_json={
            "options": [
                {"code": "ok", "label": "合格", "skip_process_ids": [processes[2].id]},
                {"code": "repair", "label": "维修", "skip_process_ids": []},
            ],
        }))
    db.flush()
    module = migration()
    module.seed_routes(db.connection(), module.route_plans(db.connection()))
    for product_type in ("cap", "piece"):
        for field, dict_type in C.ATTR_DICTS[product_type].items():
            value = attrs(product_type).get(field)
            if value:
                db.add(SysDict(type=dict_type, code=value, label=value, sort=0, is_active=True))
    db.commit()
    return user


def attrs(product_type="cap"):
    if product_type == "piece":
        return {"product_type": "piece", "craft": "全递针9*14", "length": "25厘米"}
    return {"product_type": "cap", "craft": "递旋", "size": "59", "length": "20厘米"}


def payload(product_type="cap", **changes):
    return OrderCreate.model_validate({
        "order_kind": "production", "request_id": str(uuid4()), "order_date": "2026-09-07",
        "items": [{"client_key": "line-1", "attrs": attrs(product_type), "order_qty": 2}],
        **changes,
    })


@pytest.mark.parametrize("product_type", ["cap", "piece"])
def test_create_without_customer_or_prices_and_replay(db, context, monkeypatch, product_type):
    def forbidden(*args, **kwargs):
        pytest.fail("Production order must not invoke sales pricing")
    monkeypatch.setattr(pricing_service, "lock_and_validate_order_quotes", forbidden)
    request = payload(product_type)
    first = order_service.create_order(db, request, context.id)
    replay = order_service.create_order(db, request, context.id)
    assert first["domestic_no"].startswith("DP")
    assert first["total_amount"] == 0
    assert replay["id"] == first["id"] and replay["replayed"]
    order = db.get(DomesticOrder, first["id"])
    assert order.order_no == order.domestic_no
    assert (order.customer_id, order.order_category, order.order_type, order.order_channel, order.required_ship_date) == (None,) * 5
    assert db.query(DomesticCustomer).count() == db.query(DomesticCustomerLedger).count() == 0
    item = db.query(DomesticOrderItem).filter_by(order_id=order.id).one()
    assert (item.unit_price, item.original_price, item.discount_amount, item.labor_fee) == (Decimal(0),) * 4
    assert item.pricing_rule == "production" and item.base_price_version_snapshot == 0
    detail = order_service.get_order_detail(db, order.id)
    assert detail["order_kind_label"] == "生产订单"
    assert detail["current_expected_quotes"] == []
    assert [step["process_name"] for step in detail["items"][0]["steps"]] == ["确认下单", "毛坯质检", "毛坯维修", "入库"]
    assert detail["items"][0]["steps"][1]["rule_type"] == "decision"


def test_production_discards_sales_and_hairstyle_fields():
    item = {"client_key": "line-1", "attrs": {**attrs(), "hair_style_series": "时尚"}, "order_qty": 1,
            "special_price": "100", "labor_fee": "50", "hairstyle": "短发", "style_requirement": "剪短",
            "hairstyle_images": ["old.png"], "style_images": ["old.png"], "color": "黑色", "remark": "备货"}
    request = payload(customer_id=1, customer_shop_name="test", order_type="first", order_channel="wechat",
                      required_ship_date="2026-09-08", items=[item])
    assert request.customer_id == 1 and request.required_ship_date is None
    row = request.items[0]
    assert row.attrs.hair_style_series is None
    assert row.special_price is None and row.labor_fee == 0
    assert row.hairstyle is None and row.style_requirement is None
    assert row.hairstyle_images == row.style_images == []
    assert row.color == "黑色" and row.remark == "备货"


def test_business_still_requires_customer_dimensions_hairstyle_and_price():
    with pytest.raises(ValidationError):
        payload(order_kind="business")
    with pytest.raises(ValidationError):
        payload(order_kind="business", customer_id=1, order_no="B1", required_ship_date="2026-09-08",
                order_type="first", order_channel="wechat")


def test_production_draft_append_update_and_submit_without_money(db, context, monkeypatch):
    monkeypatch.setattr(pricing_service, "lock_and_validate_order_quotes", lambda *a, **k: pytest.fail("Pricing was called"))
    order_id = order_service.create_order(db, payload(is_draft=True), context.id)["id"]
    append = OrderItemAppend(client_key="line-2", attrs=attrs("piece"), order_qty=3, request_id=str(uuid4()))
    added = order_service.add_item(db, order_id, append, context.id)
    assert order_service.add_item(db, order_id, append, context.id)["replayed"]
    order_service.update_item(db, added["id"], OrderItemUpdate(order_qty=4, remark="优先备货"), context.id)
    with pytest.raises(ValueError, match="不设置价格"):
        order_service.update_item(db, added["id"], OrderItemUpdate(unit_price=1), context.id)
    db.rollback()
    with pytest.raises(ValueError, match="不填写发型"):
        order_service.update_item(db, added["id"], OrderItemUpdate(hairstyle="短发"), context.id)
    db.rollback()
    request = DraftSubmitRequest(request_id=str(uuid4()))
    first = order_service.submit_draft(db, order_id, request, context.id)
    assert first["status"] == C.ORDER_PRODUCING and first["charged_amount"] == 0
    assert order_service.submit_draft(db, order_id, request, context.id)["replayed"]
    assert db.query(DomesticCustomerLedger).count() == 0
    with pytest.raises(ValueError, match="销售字段"):
        order_service.update_order(db, order_id, OrderUpdate(order_type="other"), context.id)


def test_storage_finishes_production_and_shipping_is_rejected(db, context):
    order_id = order_service.create_order(db, payload(), context.id)["id"]
    item = db.query(DomesticOrderItem).filter_by(order_id=order_id).one()
    rows = db.query(DomesticItemProgress).filter_by(item_id=item.id).all()
    for row in rows:
        db.add(UserProcessBinding(user_id=context.id, process_id=row.process_id))
    db.commit()
    for row in rows:
        report_service.submit_report(
            db, item_id=item.id, progress_id=row.id, qty=2, user_id=context.id,
            request_id=str(uuid4()), outcomes={"ok": 0, "repair": 2} if row.step_order == 2 else None,
        )
    assert item.status == C.ITEM_DONE and db.get(DomesticOrder, order_id).status == C.ORDER_DONE
    with pytest.raises(ValueError, match="无需登记发货"):
        order_service.ship_item(db, item.id, ItemShipRequest(ship_time="2026-09-07T12:00:00", ship_weight=100), context.id)


@pytest.mark.parametrize("kind,category,count,first,last", [
    ("production", None, 4, "确认下单", "入库"),
    ("business", "normal", 3, "毛坯出库", "发货完成"),
    ("business", "special", 7, "确认下单", "发货完成"),
])
@pytest.mark.parametrize("product_type", ["cap", "piece"])
def test_six_route_combinations(db, context, kind, category, count, first, last, product_type):
    from app.domestic.order_kind_service import resolve_order_route
    rid = resolve_order_route(db, order_kind=kind, order_category=category, product_type=product_type)
    steps = db.query(Process.name).join(ProcessRouteStep, ProcessRouteStep.process_id == Process.id).filter(
        ProcessRouteStep.route_id == rid).order_by(ProcessRouteStep.step_order).all()
    assert len(steps) == count and steps[0][0] == first and steps[-1][0] == last
    rules = route_rule_service.list_rules(db, rid)
    assert len(rules) == (0 if category == "normal" else 1)


def test_list_filter_and_number_sequences_are_separate(db, context):
    result = order_service.create_order(db, payload(), context.id)
    assert order_service.list_orders(db, order_kind="production")[1] == 1
    assert order_service.list_orders(db, order_kind="business")[1] == 0
    assert order_service._generate_domestic_no(db, "business").endswith("-001")
    assert order_service._generate_domestic_no(db, "production").endswith("-002")
    assert order_service.get_order_detail(db, result["id"])["order_kind"] == "production"


def test_same_piece_sku_keeps_three_independent_order_route_snapshots(db, context):
    from tests.test_domestic_member_pricing import _order_pricing_context, _priced_order_payload

    production_id = order_service.create_order(db, payload("piece"), context.id)["id"]
    user, customer, product, _base, expected, item_attrs = _order_pricing_context(db, "same-sku", attrs=attrs("piece"))
    normal = _priced_order_payload(customer, item_attrs, expected, request_id="same-sku-normal")
    normal_id = order_service.create_order(db, normal, user.id)["id"]
    special = _priced_order_payload(customer, item_attrs, expected, request_id="same-sku-special")
    special.order_category = "special"
    special.items[0].special_price = Decimal("800")
    special_id = order_service.create_order(db, special, user.id)["id"]
    rows = [db.query(DomesticOrderItem).filter_by(order_id=oid).one() for oid in (production_id, normal_id, special_id)]
    assert {item.product_id for item in rows} == {product.id}
    assert len({item.route_id for item in rows}) == 3
    assert [len(order_service.get_order_detail(db, oid)["items"][0]["steps"]) for oid in (production_id, normal_id, special_id)] == [4, 3, 7]
    assert rows[0].unit_price == 0 and rows[1].unit_price == expected["discount_price"] and rows[2].unit_price == 800
    assert db.query(DomesticCustomerLedger).filter_by(order_id=production_id).count() == 0


def test_number_sequence_crosses_three_digits(db, context):
    order_id = order_service.create_order(db, payload(), context.id)["id"]
    order = db.get(DomesticOrder, order_id)
    prefix = order.domestic_no.rsplit("-", 1)[0]
    order.domestic_no = f"{prefix}-999"
    db.commit()
    order_service.create_order(db, payload(), context.id)
    assert order_service._generate_domestic_no(db, "production") == f"{prefix}-1001"


def test_production_export_drops_sales_and_hairstyle_columns(db, context):
    from openpyxl import load_workbook
    from app.domestic.export_service import build_order_workbook

    order_id = order_service.create_order(db, payload(), context.id)["id"]
    sheet = load_workbook(build_order_workbook(order_service.get_order_detail(db, order_id))).active
    assert sheet.title == "内贸生产备货单"
    assert sheet.max_column == 12
    text = " ".join(str(cell.value or "") for row in sheet for cell in row)
    assert "公司毛坯备货" in text
    assert not any(field in text for field in ("客户订单号", "客户：", "要求发货", "订单类别", "订单类型", "订单渠道", "发型系列", "发型要求"))


def test_migration_rejects_missing_boundary_before_ddl(db, context):
    # Remove generated targets only in the isolated test database, then invalidate a source boundary.
    for route in db.query(ProcessRoute).filter(ProcessRoute.name.like("% · %")).all():
        db.query(DomesticRouteRule).filter_by(route_id=route.id).delete()
        db.query(ProcessRouteStep).filter_by(route_id=route.id).delete()
        db.delete(route)
    db.query(Process).filter_by(name="入库").one().name = "不可识别的终点"
    db.flush()
    with pytest.raises(RuntimeError, match="unique boundary"):
        migration().route_plans(db.connection())


def test_migration_compiles_mysql_nullable_columns_and_zero_price_constraints(monkeypatch):
    from io import StringIO
    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    output = StringIO()
    operations = Operations(MigrationContext.configure(
        dialect_name="mysql", opts={"as_sql": True, "output_buffer": output},
    ))
    module = migration()
    monkeypatch.setattr(module, "op", operations)
    monkeypatch.setattr(module, "route_plans", lambda conn: [])
    monkeypatch.setattr(module, "seed_routes", lambda conn, plans: None)
    module.upgrade()
    sql = output.getvalue()
    assert "ADD COLUMN order_kind VARCHAR(16) NOT NULL" in sql
    assert "DEFAULT 'business'" in sql
    assert "MODIFY customer_id INTEGER NULL" in sql
    assert "MODIFY order_category VARCHAR(16) NULL" in sql
    assert "ck_dom_order_kind_fields" in sql and "customer_id IS NULL" in sql
    assert "ck_dom_item_production_price" in sql and "base_price_version_snapshot = 0" in sql


@pytest.mark.parametrize("draft", [False, True])
def test_production_can_select_customer_without_sales_side_effects(db, context, monkeypatch, draft):
    customer = DomesticCustomer(shop_name="生产关联客户", balance=Decimal("1234.00"), created_by=context.id,
                                last_order_date=date(2026, 8, 1))
    db.add(customer)
    db.commit()
    monkeypatch.setattr(pricing_service, "lock_and_validate_order_quotes", lambda *a, **k: pytest.fail("Pricing was called"))
    request = payload(customer_id=customer.id, is_draft=draft)
    result = order_service.create_order(db, request, context.id)
    detail = order_service.get_order_detail(db, result["id"])
    assert detail["customer_id"] == customer.id
    assert detail["customer_name"] == "生产关联客户"
    from app.domestic.router import get_item_unit_qrcodes
    from app.domestic.export_service import build_order_workbook
    from openpyxl import load_workbook
    labels = get_item_unit_qrcodes(detail["items"][0]["id"], start_no=1, end_no=1, db=db,
                                   _user={"sub": str(context.id), "permissions": ["domestic:read"], "roles": []})["data"]
    assert labels["customer_name"] == "生产关联客户" and labels["order_date"] == date(2026, 9, 7)
    sheet = load_workbook(build_order_workbook(detail)).active
    assert "生产关联客户" in sheet["A3"].value
    assert detail["total_amount"] == detail["charged_amount"] == 0
    if draft:
        order_service.submit_draft(db, result["id"], DraftSubmitRequest(request_id=str(uuid4())), context.id)
    db.refresh(customer)
    assert customer.balance == Decimal("1234.00")
    assert customer.last_order_date == date(2026, 8, 1)
    assert db.query(DomesticCustomerLedger).count() == 0
    assert order_service.create_order(db, request, context.id)["replayed"] is True


def test_production_edit_customer_can_select_clear_and_reject_missing(db, context):
    customer = DomesticCustomer(shop_name="可选生产客户", balance=Decimal("100"), created_by=context.id)
    db.add(customer)
    db.commit()
    result = order_service.create_order(db, payload(), context.id)
    order_service.update_order(db, result["id"], OrderUpdate(production_customer_id=customer.id), context.id)
    assert order_service.get_order_detail(db, result["id"])["customer_id"] == customer.id
    order_service.update_order(db, result["id"], OrderUpdate(production_customer_id=None), context.id)
    assert order_service.get_order_detail(db, result["id"])["customer_id"] is None
    with pytest.raises(ValueError, match="客户不存在"):
        order_service.update_order(db, result["id"], OrderUpdate(production_customer_id=999999), context.id)
    assert db.query(DomesticCustomerLedger).count() == 0
