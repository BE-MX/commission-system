"""Confirmed owner bindings for both print formats, using isolated SQLite."""
import io

from docx import Document
from app.auth.models import ArkUserExternalBinding
from app.shipping_inspection.print_service import with_owner_chinese_name
from tests.test_shipping_inspection import _user, _pc_client


def bind(db, user, name="Alice", ident=1, **kwargs):
    db.add(ArkUserExternalBinding(id=ident, ark_user_id=user.id, provider="okki",
        external_account_id=str(ident), external_display_name=name,
        binding_status="active", **kwargs))
    db.flush()


def test_owner_name_matching_and_ambiguity(db):
    user = _user(db)
    user.real_name = "王小红"
    bind(db, user)
    record = {"owner_name": " alice "}
    assert with_owner_chinese_name(db, record)["owner_name"] == "alice（王小红）"
    assert record["owner_name"] == " alice "
    for value in (None, "Unknown", "王小红", "Alice（王小红）"):
        assert with_owner_chinese_name(db, {"owner_name": value})["owner_name"] == value
    bind(db, user, ident=2)  # Duplicate bindings to the same user are unambiguous.
    assert with_owner_chinese_name(db, record)["owner_name"] == "alice（王小红）"
    other = _user(db, "other")
    other.real_name = "李小明"
    bind(db, other, ident=3)
    assert with_owner_chinese_name(db, record) == record


def test_inactive_deleted_or_non_chinese_binding_is_not_used(db):
    user = _user(db)
    bind(db, user)
    record = {"owner_name": "Alice"}
    assert with_owner_chinese_name(db, record) == record
    user.real_name = "王小红"
    binding = db.query(ArkUserExternalBinding).first()
    binding.binding_status = "inactive"
    db.flush()
    assert with_owner_chinese_name(db, record) == record
    binding.binding_status = "active"
    from app.core.time import beijing_now
    binding.deleted_at = beijing_now()
    db.flush()
    assert with_owner_chinese_name(db, record) == record


def test_html_payload_and_word_show_same_owner(db, monkeypatch):
    user = _user(db, "Alice")
    user.real_name = "王小红"
    bind(db, user, name="王小红")
    from app.shipping_inspection import outbound_service
    monkeypatch.setattr(outbound_service, "get_outbound_record", lambda *a, **k: {
        "owner_name": "Alice", "outbound_no": "TEST"})
    monkeypatch.setattr(outbound_service, "list_outbound_items", lambda *a: [])
    with _pc_client(db, user, [], roles=["super_admin"]) as client:
        data = client.get("/api/shipping-inspection/outbound-records/TEST/print-data")
        assert data.status_code == 200
        assert data.json()["data"]["record"]["owner_name"] == "Alice（王小红）"
        word = client.get("/api/shipping-inspection/outbound-records/TEST/word")
        assert word.status_code == 200
        document = Document(io.BytesIO(word.content))
        assert document.tables[0].cell(2, 1).text == "Alice（王小红）"


def test_username_matches_when_okki_display_name_is_chinese(db):
    user = _user(db, "Ginny")
    user.real_name = "翟佳盟"
    bind(db, user, name="翟佳盟")
    assert with_owner_chinese_name(db, {"owner_name": " ginny "})["owner_name"] == "ginny（翟佳盟）"


def test_username_without_binding_and_cross_field_conflict(db):
    user = _user(db, "Alice")
    user.real_name = "王小红"
    record = {"owner_name": "Alice"}
    assert with_owner_chinese_name(db, record)["owner_name"] == "Alice（王小红）"
    other = _user(db, "Other")
    other.real_name = "李小明"
    bind(db, other, name="Alice")
    assert with_owner_chinese_name(db, record) == record


def test_deleted_username_is_not_used(db):
    from app.core.time import beijing_now
    user = _user(db, "Alice")
    user.real_name = "王小红"
    user.deleted_at = beijing_now()
    db.flush()
    record = {"owner_name": "Alice"}
    assert with_owner_chinese_name(db, record) == record


def test_print_owner_uses_live_handler_not_api_creator(db, monkeypatch):
    from app.invoice import okki_client
    monkeypatch.setattr(okki_client, "get_outbound_info", lambda db, ident: {
        "outbound_invoice_id": 123, "handler_info": [{"user_id": "57125949", "nickname": "Eva"}],
        "create_user_info": {"nickname": "Rainy"},
    })
    record = {"outbound_invoice_id": "123", "owner_name": "Rainy"}
    assert with_owner_chinese_name(db, record)["owner_name"] == "Eva"
    assert record["owner_name"] == "Rainy"


def test_print_owner_lookup_failure_does_not_print_creator(db, monkeypatch):
    import pytest
    from fastapi import HTTPException
    from app.invoice import okki_client
    def fail(*args):
        raise okki_client.OkkiApiError("timeout")
    monkeypatch.setattr(okki_client, "get_outbound_info", fail)
    with pytest.raises(HTTPException) as exc:
        with_owner_chinese_name(db, {"outbound_invoice_id": "123", "owner_name": "Rainy"})
    assert exc.value.status_code == 502


def test_print_owner_multiple_handlers_and_no_handler(db, monkeypatch):
    from app.invoice import okki_client
    monkeypatch.setattr(okki_client, "get_outbound_info", lambda *args: {
        "outbound_invoice_id": 123, "handler_info": [{"nickname": "Eva"}, {"nickname": "Mary"}, {"nickname": "Eva"}],
    })
    record = {"outbound_invoice_id": "123", "owner_name": "Rainy"}
    assert with_owner_chinese_name(db, record)["owner_name"] == "Eva / Mary"
    monkeypatch.setattr(okki_client, "get_outbound_info", lambda *args: {"outbound_invoice_id": 123, "handler_info": []})
    assert with_owner_chinese_name(db, record)["owner_name"] is None


def test_html_and_word_use_live_outbound_handler(db, monkeypatch):
    from app.invoice import okki_client
    from app.shipping_inspection import outbound_service
    user = _user(db, "Admin")
    monkeypatch.setattr(okki_client, "get_outbound_info", lambda *args: {
        "outbound_invoice_id": 123, "handler_info": [{"nickname": "Eva"}],
    })
    monkeypatch.setattr(outbound_service, "get_outbound_record", lambda *a, **k: {
        "outbound_invoice_id": "123", "owner_name": "Rainy", "outbound_no": "ly914首返出库单"})
    monkeypatch.setattr(outbound_service, "list_outbound_items", lambda *a: [])
    with _pc_client(db, user, [], roles=["super_admin"]) as client:
        response = client.get("/api/shipping-inspection/outbound-records/123/print-data")
        assert response.status_code == 200
        assert response.json()["data"]["record"]["owner_name"] == "Eva"
        response = client.get("/api/shipping-inspection/outbound-records/123/word")
        assert response.status_code == 200
        document = Document(io.BytesIO(response.content))
        assert document.tables[0].cell(2, 1).text == "Eva"


def test_print_persists_refreshed_okki_token(db, monkeypatch):
    from sqlalchemy.orm import Session
    from app.invoice import okki_client
    from app.invoice.models import XiaomanSettings
    row = XiaomanSettings(id=1, access_token="expired-test-token")
    db.add(row)
    db.commit()
    def refreshed(session, ident):
        session.get(XiaomanSettings, 1).access_token = "fresh-test-token"
        return {"outbound_invoice_id": 123, "handler_info": [{"nickname": "Eva"}]}
    monkeypatch.setattr(okki_client, "get_outbound_info", refreshed)
    assert with_owner_chinese_name(db, {"outbound_invoice_id": "123", "owner_name": "Rainy"})["owner_name"] == "Eva"
    db.rollback()  # Request end must not undo the refreshed token.
    with Session(db.bind) as followup:
        assert followup.get(XiaomanSettings, 1).access_token == "fresh-test-token"
