"""薪资 M4 导出：工资表双 sheet / 银行代发盘 / 导出留痕与作废。

每条测试对应一个真实的出错形态：
- 结构：财务拿导出件直接替换手工表，表头、分组、合计、签字栏任何一处走样都要返工；
- 快照：锁定后改档案（换卡、调岗）不能改写已锁定月份的导出——§2.5 错误 2 的系统版；
- 作废：解锁前发出去的文件必须能被认出来是废的；
- 明文：卡号明文只在有权限时出现在文件里，永不进事件 payload 与日志。
"""

import csv
import datetime as dt
import io
import json
from decimal import Decimal

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.auth.dependencies import get_current_user
from app.auth.models import ArkUser
from app.core.config import get_settings
from app.core.database import Base, get_db
from app.salary import calc_service as cs
from app.salary import export_service as es
from app.salary import period_service as ps
from app.salary import router as salary_router
from app.salary.models import (
    SalaryAttendance,
    SalaryChangeLog,
    SalaryDeptMapping,
    SalaryEmployeeProfile,
    SalaryFundImport,
    SalaryGradeTable,
    SalaryInsuranceImport,
    SalaryPeriod,
    SalaryPeriodEvent,
    SalaryRecord,
    SalaryRuleParam,
)
from app.salary.pii import encrypt_pii
from app.salary.seed import seed_rule_params

D = Decimal

_TEST_ENC_KEY = "dGVzdC1zYWxhcnktZW5jLWtleS0zMi1ieXRlcy0hIQ=="
_TEST_HASH_KEY = "test-salary-hash-key"

# 测试卡号（非真实卡）。19 位：数字格式下 Excel 会丢末几位，用来验证写成文本
CARDS = {
    "1": "6217000000000000011",
    "2": "6217000000000000022",
    "3": "6217000000000000033",
    "4": "6217000000000000044",
}


@pytest.fixture(autouse=True)
def salary_keys(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "ARK_SALARY_ENCRYPTION_KEY", _TEST_ENC_KEY)
    monkeypatch.setattr(settings, "ARK_SALARY_HASH_KEY", _TEST_HASH_KEY)


@pytest.fixture()
def db():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(
        engine,
        tables=[
            SalaryEmployeeProfile.__table__, SalaryPeriod.__table__,
            SalaryPeriodEvent.__table__, SalaryAttendance.__table__,
            SalaryInsuranceImport.__table__, SalaryFundImport.__table__,
            SalaryRecord.__table__, SalaryChangeLog.__table__,
            SalaryRuleParam.__table__, SalaryGradeTable.__table__,
            SalaryDeptMapping.__table__, ArkUser.__table__,
        ],
    )
    session = sessionmaker(bind=engine, autoflush=False)()
    seed_rule_params(session)
    session.flush()
    session.query(SalaryRuleParam).update(
        {SalaryRuleParam.effective_from: dt.date(2026, 1, 1)})
    session.add(ArkUser(id=7, username="caiwu", real_name="财务小王",
                        password_hash="x"))
    session.commit()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _profile(db, name, emp_no, **kw) -> SalaryEmployeeProfile:
    defaults = dict(status="active", payroll_included=1, fund_included=0,
                    special_calc=0, base_salary_override=D("5000"),
                    dingtalk_userid=f"U{emp_no}",
                    bank_card_cipher=encrypt_pii(CARDS.get(emp_no)),
                    bank_name="中国银行")
    defaults.update(kw)
    p = SalaryEmployeeProfile(name=name, emp_no=emp_no, **defaults)
    db.add(p)
    db.commit()
    return p


def _calculated_period(db) -> tuple[SalaryPeriod, dict[str, SalaryEmployeeProfile]]:
    """2026-03，两组两主体：后综部（行政部）2 人 + 业务部（业务一部）2 人。

    业务部的人插在前面建档，验证汇总表按映射 sort_order 而不是建档顺序分组。
    """
    db.add_all([
        SalaryDeptMapping(dept_detail="行政部", dept_group="后综部", sort_order=1),
        SalaryDeptMapping(dept_detail="业务一部", dept_group="业务部", sort_order=2),
    ])
    db.commit()
    period = ps.create_period(db, "2026-03", workday_count=22)
    for status in (ps.STATUS_ATTENDANCE, ps.STATUS_IMPORTED):
        period = ps.transition(db, period, status, expected_version=period.status_version)
    people = {
        "吕德洋": _profile(db, "吕德洋", "1", dept_detail="业务一部", position="业务员",
                         hire_date=dt.date(2020, 5, 6), insurance_entity="莱莎贸易"),
        "孙正华": _profile(db, "孙正华", "2", dept_detail="行政部", position="行政",
                         hire_date=dt.date(2019, 3, 1), insurance_entity="莱莎发制品"),
        "李晓雨": _profile(db, "李晓雨", "3", dept_detail="行政部", position="出纳",
                         hire_date=dt.date(2024, 7, 15), insurance_entity="莱莎发制品"),
        "刘也": _profile(db, "刘也", "4", dept_detail="业务一部", position="业务员",
                       hire_date=dt.date(2025, 11, 3), insurance_entity="莱莎贸易"),
    }
    for p in people.values():
        db.add(SalaryAttendance(period_id=period.id, employee_id=p.id,
                                due_days=D("31"), actual_days=D("31"), full_attendance=1,
                                personal_leave_hours=D("0"), sick_leave_hours=D("0")))
    db.commit()
    cs.calculate_period(db, period, expected_version=period.status_version)
    return period, people


def _confirmed_period(db):
    period, people = _calculated_period(db)
    period = ps.transition(db, period, ps.STATUS_REVIEWING,
                           expected_version=period.status_version)
    period = ps.confirm(db, period, expected_version=period.status_version, operator_id=7)
    return period, people


def _sheet_rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _exports(db, period_id) -> list[SalaryPeriodEvent]:
    return db.query(SalaryPeriodEvent).filter_by(
        period_id=period_id, event_type="export").order_by(SalaryPeriodEvent.id).all()


# ---------------------------------------------------------------------------
# 工资表结构：与 3 月人工表同构
# ---------------------------------------------------------------------------

def test_workbook_matches_manual_sheet_structure(db):
    period, people = _confirmed_period(db)
    filename, content = es.export_workbook(db, period, plain_card=False, operator_id=7)
    assert filename == "2026-03工资表_第1次导出.xlsx"

    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["2026年3月工资汇总表", "2026年3月工资表（明细）"]

    ws1 = wb.worksheets[0]
    rows = _sheet_rows(ws1)
    assert rows[0][0] == "2026年3月工资汇总表"
    assert rows[1] == es.SUMMARY_HEADERS
    # 后综部（sort 1）在前，组内序号从 1 重排，部门名只在组内首行、整列纵向合并
    assert [r[:3] for r in rows[2:4]] == [[1, "后综部", "孙正华"], [2, None, "李晓雨"]]
    assert rows[4][0] == "合计"
    assert [r[:3] for r in rows[5:7]] == [[1, "业务部", "吕德洋"], [2, None, "刘也"]]
    assert rows[7][0] == "合计"
    assert rows[8][0] == "总计"
    merged = {str(m) for m in ws1.merged_cells.ranges}
    assert {"A1:I1", "B3:B4", "B6:B7", "A5:G5", "A8:G8", "A9:G9"} <= merged

    net = {r.employee_id: r.net_salary for r in db.query(SalaryRecord).filter_by(period_id=period.id)}
    hz = net[people["孙正华"].id] + net[people["李晓雨"].id]
    yw = net[people["吕德洋"].id] + net[people["刘也"].id]
    assert D(str(rows[4][7])) == hz
    assert D(str(rows[7][7])) == yw
    assert D(str(rows[8][7])) == hz + yw

    # 入职日期是日期单元格；卡号脱敏且是文本
    assert rows[2][4] == dt.datetime(2019, 3, 1)
    assert ws1["E3"].number_format == "yyyy/mm/dd"
    assert rows[2][5] == "6217***********0022"
    assert ws1["F3"].data_type == "s"
    assert CARDS["2"] not in json.dumps(rows, default=str)

    # 签字栏 + 导出编号脚注
    flat = [c for r in rows[9:] for c in r if c]
    assert {"制表：", "审核：", "复核："} <= set(flat)
    assert any(str(c).startswith("方舟系统导出 第1次") and "财务小王" in str(c) for c in flat)

    ws2 = wb.worksheets[1]
    rows2 = _sheet_rows(ws2)
    assert rows2[0] == es.DETAIL_HEADERS
    assert len(rows2[0]) == 23
    detail = rows2[1:5]
    # 明细按 seq_no（不按分组），姓名在第 4 列
    assert [r[0] for r in detail] == [1, 2, 3, 4]
    total_row = rows2[5]
    assert total_row[0] == "合计"
    assert D(str(total_row[20])) == hz + yw                       # 实发工资
    for col in range(8, 23):
        assert D(str(total_row[col] or 0)) == sum(
            (D(str(r[col] or 0)) for r in detail), D("0")), es.DETAIL_HEADERS[col]
    flat2 = [c for r in rows2[6:] for c in r if c]
    assert {"制表：", "审核：", "复核："} <= set(flat2)


def test_plain_card_workbook_writes_full_card_as_text(db):
    period, _ = _confirmed_period(db)
    _, content = es.export_workbook(db, period, plain_card=True, operator_id=7)
    ws1 = load_workbook(io.BytesIO(content)).worksheets[0]
    assert ws1["F3"].value == CARDS["2"]
    assert ws1["F3"].data_type == "s"   # 19 位不被 Excel 变成科学计数


# ---------------------------------------------------------------------------
# 快照：锁定批次只读 snap_*
# ---------------------------------------------------------------------------

def test_confirmed_export_ignores_profile_changes_after_lock(db):
    period, people = _confirmed_period(db)
    p = people["孙正华"]
    p.name = "孙改名"
    p.bank_card_cipher = encrypt_pii("6222999999999999999")
    p.hire_date = dt.date(2025, 1, 1)
    p.insurance_entity = "莱莎贸易"
    p.position = "经理"
    db.commit()

    _, content = es.export_workbook(db, period, plain_card=True, operator_id=7)
    row = _sheet_rows(load_workbook(io.BytesIO(content)).worksheets[0])[2]
    assert row[2:6] == ["孙正华", "行政", dt.datetime(2019, 3, 1), CARDS["2"]]

    # 银行盘也按锁定时的主体与卡号
    _, csv_bytes = es.export_bank(db, period, entity="莱莎发制品", operator_id=7)
    text = csv_bytes.decode("utf-8-sig")
    assert CARDS["2"] in text and "6222999999999999999" not in text


def test_confirm_freezes_hire_date_and_insurance_entity(db):
    period, people = _confirmed_period(db)
    row = db.query(SalaryRecord).filter_by(
        period_id=period.id, employee_id=people["吕德洋"].id).one()
    assert row.snap_hire_date == dt.date(2020, 5, 6)
    assert row.snap_insurance_entity == "莱莎贸易"


def test_control_characters_in_text_do_not_break_export(db):
    """档案里粘进来的不可见控制字符（ 等）openpyxl 会直接抛错，导出前剔除。"""
    period, people = _confirmed_period(db)
    db.query(SalaryRecord).filter_by(
        period_id=period.id, employee_id=people["孙正华"].id).update(
        {SalaryRecord.remark_summary: "补发三月交通费"})
    db.commit()
    _, content = es.export_workbook(db, period, plain_card=False, operator_id=7)
    flat = [c for r in _sheet_rows(load_workbook(io.BytesIO(content)).worksheets[0]) for c in r if c]
    assert "补发三月交通费" in flat


def test_missing_card_is_flagged_not_left_blank(db):
    period, people = _confirmed_period(db)
    db.query(SalaryRecord).filter_by(
        period_id=period.id, employee_id=people["刘也"].id).update(
        {SalaryRecord.snap_bank_card_cipher: None})
    db.commit()
    assert es.export_overview(db, period)["workbook"]["missing_card_names"] == ["刘也"]
    _, content = es.export_workbook(db, period, plain_card=True, operator_id=7)
    rows = _sheet_rows(load_workbook(io.BytesIO(content)).worksheets[0])
    liu = next(r for r in rows if r[2] == "刘也")
    assert liu[5] == es.CARD_MISSING
    flat = [str(c) for r in rows for c in r if c]
    assert any("1 人无银行卡号" in c for c in flat)


def test_overview_explains_unfrozen_locked_batch(db):
    period, _ = _confirmed_period(db)
    db.query(SalaryRecord).filter_by(period_id=period.id).update(
        {SalaryRecord.snapshot_at: None})
    db.commit()
    wb = es.export_overview(db, period)["workbook"]
    assert wb["available"] is False and "解锁后重新锁定" in wb["reason"]


def test_locked_batch_without_snapshot_is_refused_not_backfilled(db):
    """锁定早于快照机制的批次：拒导，不回落活档案。"""
    period, _ = _confirmed_period(db)
    db.query(SalaryRecord).filter_by(period_id=period.id).update(
        {SalaryRecord.snapshot_at: None})
    db.commit()
    with pytest.raises(es.SalaryExportError, match="解锁后重新锁定"):
        es.export_workbook(db, period, plain_card=False, operator_id=7)
    assert _exports(db, period.id) == []


# ---------------------------------------------------------------------------
# 预览稿
# ---------------------------------------------------------------------------

def test_unlocked_batch_exports_marked_preview(db):
    period, _ = _calculated_period(db)
    filename, content = es.export_workbook(db, period, plain_card=False, operator_id=7)
    assert "预览稿" in filename
    ws1 = load_workbook(io.BytesIO(content)).worksheets[0]
    assert "不得作为发放依据" in ws1["A1"].value
    assert "预览稿" in ws1.oddHeader.center.text
    assert es.export_history(db, period)[0]["status"] == "preview"


def test_draft_batch_has_nothing_to_export(db):
    period = ps.create_period(db, "2026-03", workday_count=22)
    with pytest.raises(es.SalaryExportError, match="还没有计算"):
        es.export_workbook(db, period, plain_card=False, operator_id=7)


# ---------------------------------------------------------------------------
# 银行代发盘
# ---------------------------------------------------------------------------

def _csv_rows(content: bytes) -> list[list[str]]:
    assert content.startswith(b"\xef\xbb\xbf")  # BOM：Excel 打开中文不乱码
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


def test_bank_csv_split_by_insurance_entity(db):
    period, _ = _confirmed_period(db)
    overview = es.bank_overview(db, period)
    assert overview["available"] is True
    assert [g["entity"] for g in overview["groups"]] == ["莱莎发制品", "莱莎贸易"]
    assert all(g["count"] == 2 and not g["issues"] for g in overview["groups"])

    name, content = es.export_bank(db, period, entity="莱莎发制品", operator_id=7)
    assert name == "2026-03银行代发_莱莎发制品_第1次导出.csv"
    rows = _csv_rows(content)
    assert rows[0] == ["姓名", "卡号", "金额"]
    assert {r[0] for r in rows[1:]} == {"孙正华", "李晓雨"}
    assert {r[1] for r in rows[1:]} == {CARDS["2"], CARDS["3"]}
    recs = {r.snap_name: r for r in db.query(SalaryRecord).filter_by(period_id=period.id)}
    assert rows[1][2] == f"{recs[rows[1][0]].net_after_tax:.2f}"


def test_bank_csv_single_file_when_entity_omitted(db):
    period, _ = _confirmed_period(db)
    name, content = es.export_bank(db, period, entity=None, operator_id=7)
    assert "_全部_" in name
    assert len(_csv_rows(content)) == 5


def test_bank_blocks_missing_card_and_records_nothing(db):
    period, people = _confirmed_period(db)
    db.query(SalaryRecord).filter_by(
        period_id=period.id, employee_id=people["刘也"].id).update(
        {SalaryRecord.snap_bank_card_cipher: None})
    db.commit()
    with pytest.raises(es.SalaryExportError, match="刘也"):
        es.export_bank(db, period, entity="莱莎贸易", operator_id=7)
    assert _exports(db, period.id) == []
    # 另一主体不受影响
    es.export_bank(db, period, entity="莱莎发制品", operator_id=7)


def test_bank_blocks_negative_and_skips_zero(db):
    period, people = _confirmed_period(db)
    db.query(SalaryRecord).filter_by(
        period_id=period.id, employee_id=people["刘也"].id).update(
        {SalaryRecord.net_after_tax: D("0")})
    db.commit()
    group = [g for g in es.bank_overview(db, period)["groups"] if g["entity"] == "莱莎贸易"][0]
    assert group["count"] == 1 and group["zero_names"] == ["刘也"]
    _, content = es.export_bank(db, period, entity="莱莎贸易", operator_id=7)
    assert [r[0] for r in _csv_rows(content)[1:]] == ["吕德洋"]

    db.query(SalaryRecord).filter_by(
        period_id=period.id, employee_id=people["刘也"].id).update(
        {SalaryRecord.net_after_tax: D("-5")})
    db.commit()
    with pytest.raises(es.SalaryExportError, match="负数"):
        es.export_bank(db, period, entity="莱莎贸易", operator_id=7)


def test_bank_requires_locked_batch(db):
    period, _ = _calculated_period(db)
    with pytest.raises(es.SalaryExportError, match="锁定后"):
        es.export_bank(db, period, entity=None, operator_id=7)
    assert es.bank_overview(db, period)["available"] is False


# ---------------------------------------------------------------------------
# 留痕与作废（决策 A4）
# ---------------------------------------------------------------------------

def test_unlock_voids_earlier_exports_and_new_file_says_so(db):
    period, _ = _confirmed_period(db)
    es.export_workbook(db, period, plain_card=False, operator_id=7)
    es.export_bank(db, period, entity="莱莎贸易", operator_id=7)

    period = ps.unlock(db, period, "孙正华银行卡录错", expected_version=period.status_version)
    period = ps.confirm(db, period, expected_version=period.status_version, operator_id=7)
    _, content = es.export_workbook(db, period, plain_card=False, operator_id=7)

    history = es.export_history(db, period)
    assert [(h["export_no"], h["status"]) for h in history] == [
        (3, "valid"), (2, "voided"), (1, "voided")]
    assert history[1]["void_reason"] == "孙正华银行卡录错"
    assert history[1]["voided_at"]
    assert history[0]["operator_name"] == "财务小王"

    flat = [c for r in _sheet_rows(load_workbook(io.BytesIO(content)).worksheets[0]) for c in r if c]
    stamp = next(str(c) for c in flat if str(c).startswith("方舟系统导出"))
    assert "第3次" in stamp and "替代第 1 次导出" in stamp


def test_void_without_unlock_event_says_trace_missing(db):
    """unlock 留痕走 log_event（失败不挡解锁）：缺事件时仍按版本判作废，原因如实说缺。"""
    period, _ = _confirmed_period(db)
    es.export_workbook(db, period, plain_card=False, operator_id=7)
    period = ps.unlock(db, period, "改档案", expected_version=period.status_version)
    db.query(SalaryPeriodEvent).filter_by(period_id=period.id, event_type="unlock").delete()
    db.commit()
    h = es.export_history(db, period)[0]
    assert h["status"] == "voided"
    assert h["voided_at"] is None and "解锁留痕缺失" in h["void_reason"]


def test_void_is_decided_by_version_not_clock(db):
    """导出记在解锁前的版本上就作废，哪怕 created_at 因时钟偏差晚于解锁时间。"""
    period, _ = _confirmed_period(db)
    es.export_workbook(db, period, plain_card=False, operator_id=7)
    period = ps.unlock(db, period, "改档案", expected_version=period.status_version)
    exp = _exports(db, period.id)[0]
    exp.created_at = period.unlocked_at + dt.timedelta(hours=1)
    db.commit()
    assert es.export_history(db, period)[0]["status"] == "voided"


def test_export_no_is_sequential_across_kinds(db):
    period, _ = _confirmed_period(db)
    es.export_workbook(db, period, plain_card=False, operator_id=7)
    es.export_bank(db, period, entity="莱莎贸易", operator_id=7)
    name, _ = es.export_workbook(db, period, plain_card=False, operator_id=7)
    assert [e.payload["export_no"] for e in _exports(db, period.id)] == [1, 2, 3]
    assert name.endswith("第3次导出.xlsx")
    db.refresh(period)
    assert period.export_seq == 3


def test_state_change_between_read_and_write_refuses_export(db, monkeypatch):
    """取数之后、落留痕之前有人解锁：拒导，不出文件、不留痕（防锁定前数据配锁定后编号）。"""
    period, _ = _confirmed_period(db)
    real_load = es.load_export_rows

    def load_then_unlock(session, p):
        rows = real_load(session, p)
        ps.unlock(session, p, "并发解锁", expected_version=p.status_version)
        return rows

    monkeypatch.setattr(es, "load_export_rows", load_then_unlock)
    with pytest.raises(es.SalaryExportError, match="刚刚被他人变更"):
        es.export_workbook(db, period, plain_card=False, operator_id=7)
    assert _exports(db, period.id) == []


def test_commit_failure_records_nothing(db, monkeypatch):
    period, _ = _confirmed_period(db)

    def fail_commit():
        raise RuntimeError("connection lost")

    monkeypatch.setattr(db, "commit", fail_commit)
    with pytest.raises(RuntimeError):
        es.export_workbook(db, period, plain_card=False, operator_id=7)
    monkeypatch.undo()
    assert _exports(db, period.id) == []


def test_export_event_and_logs_never_contain_card_numbers(db, capsys):
    period, _ = _confirmed_period(db)
    es.export_workbook(db, period, plain_card=True, operator_id=7)
    es.export_bank(db, period, entity=None, operator_id=7)
    events = _exports(db, period.id)
    assert len(events) == 2
    dumped = json.dumps([e.payload for e in events], ensure_ascii=False)
    out = capsys.readouterr().out
    for card in CARDS.values():
        assert card not in dumped
        assert card not in out
    assert events[0].payload["plain_card"] is True
    assert len(events[0].payload["sha256"]) == 64
    assert events[1].payload["record_count"] == 4


def test_build_failure_returns_no_file_and_records_nothing(db, monkeypatch):
    period, _ = _confirmed_period(db)

    def boom(*a, **kw):
        raise RuntimeError("disk full")

    monkeypatch.setattr(es, "build_workbook", boom)
    with pytest.raises(RuntimeError):
        es.export_workbook(db, period, plain_card=False, operator_id=7)
    assert _exports(db, period.id) == []


def test_stamp_uses_beijing_time_across_midnight(db, monkeypatch):
    """导出时间走 beijing_now：北京 00:05 跨日时脚注写当天日期（服务器时区无关）。"""
    period, _ = _confirmed_period(db)
    monkeypatch.setattr(es, "beijing_now", lambda: dt.datetime(2026, 4, 1, 0, 5))
    _, content = es.export_workbook(db, period, plain_card=False, operator_id=7)
    flat = [c for r in _sheet_rows(load_workbook(io.BytesIO(content)).worksheets[0]) for c in r if c]
    assert any("2026-04-01 00:05" in str(c) for c in flat)
    assert _exports(db, period.id)[0].created_at == dt.datetime(2026, 4, 1, 0, 5)


# ---------------------------------------------------------------------------
# 路由：权限与信封
# ---------------------------------------------------------------------------

def _client(db, permissions, roles=()):
    app = FastAPI()
    app.include_router(salary_router.router, prefix="/api/salary")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: {
        "sub": "7", "roles": list(roles), "permissions": list(permissions),
    }
    return TestClient(app)


def test_router_plain_card_requires_admin(db):
    period, _ = _confirmed_period(db)
    url = f"/api/salary/periods/{period.id}/export/workbook"
    assert _client(db, ["salary:read"]).get(url, params={"plain_card": True}).status_code == 403
    resp = _client(db, ["salary:read"]).get(url)
    assert resp.status_code == 200
    assert resp.headers["content-disposition"].startswith("attachment; filename*=UTF-8''")
    assert resp.headers["cache-control"] == "no-store"
    resp = _client(db, ["salary:admin"]).get(url, params={"plain_card": True})
    assert resp.status_code == 200
    ws1 = load_workbook(io.BytesIO(resp.content)).worksheets[0]
    assert ws1["F3"].value == CARDS["2"]


def test_router_bank_requires_admin(db):
    period, _ = _confirmed_period(db)
    url = f"/api/salary/periods/{period.id}/export/bank"
    assert _client(db, ["salary:write"]).get(url).status_code == 403
    resp = _client(db, [], roles=["super_admin"]).get(url, params={"entity": "莱莎贸易"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")


def test_router_export_error_is_400_with_message(db):
    period, _ = _calculated_period(db)
    resp = _client(db, ["salary:admin"]).get(f"/api/salary/periods/{period.id}/export/bank")
    assert resp.status_code == 400
    assert "锁定后" in resp.json()["detail"]


def test_router_overview_envelope(db):
    period, _ = _confirmed_period(db)
    es.export_workbook(db, period, plain_card=False, operator_id=7)
    body = _client(db, ["salary:read"]).get(f"/api/salary/periods/{period.id}/exports").json()
    assert body["code"] == 200
    data = body["data"]
    assert data["locked"] is True and data["can_plain_card"] is False
    assert data["workbook"]["available"] is True
    assert len(data["bank"]["groups"]) == 2
    assert data["history"][0]["status"] == "valid"
