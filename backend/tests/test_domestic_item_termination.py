"""Line termination preserves history and settles only the cancelled line."""
from decimal import Decimal
import pytest
from app.domestic import constants as C, order_service, balance_service, progress_service, report_service
from app.domestic.models import DomesticCustomerLedger, DomesticOrderItem
from app.domestic.schemas import OrderItemUpdate
from tests.test_domestic_item_delete import saved_order


def test_single_line_refund_and_replay_do_not_affect_sibling(db):
    user, customer, order, items = saved_order(db)
    order_service.terminate_item(db, items[0].id, '客户取消此产品', user.id)
    db.expire_all()
    assert items[0].status == C.ITEM_TERMINATED
    assert items[1].status == C.ITEM_PRODUCING
    assert order.status == C.ORDER_PRODUCING
    assert order.total_amount == order.charged_amount == Decimal('100')
    assert customer.balance == Decimal('200')
    assert '客户取消此产品' in items[0].remark
    assert db.query(DomesticOrderItem).count() == 2
    with pytest.raises(ValueError, match='已终止'):
        order_service.terminate_item(db, items[0].id, 'retry', user.id)
    assert db.query(DomesticCustomerLedger).count() == 1
    balance_service.sync_order_finance(db, order, user_id=user.id, reason='later edit')
    assert order.charged_amount == Decimal('100')
    order_service.terminate_item(db, items[1].id, '客户取消剩余产品', user.id)
    assert order.status == C.ORDER_TERMINATED
    assert order.charged_amount == order.total_amount == 0
    assert customer.balance == Decimal('300')
    assert db.query(DomesticCustomerLedger).count() == 2


@pytest.mark.parametrize('status', [0, 3, 4, 5, 6])
def test_frozen_parent_cannot_terminate_line(db, status):
    user, customer, order, items = saved_order(db, status=status)
    with pytest.raises(ValueError):
        order_service.terminate_item(db, items[0].id, '取消', user.id)
    assert items[0].status == 0 and customer.balance == 100
    assert db.query(DomesticCustomerLedger).count() == 0


def test_authorization_and_empty_reason(db):
    user, customer, order, items = saved_order(db)
    for actor, reason in [(user.id + 1, '取消'), (user.id, '  ')]:
        with pytest.raises(ValueError):
            order_service.terminate_item(db, items[0].id, reason, actor)
    assert customer.balance == 100


def test_terminated_line_cannot_edit_delete_or_report_and_stays_terminated(db):
    user, customer, order, items = saved_order(db)
    order_service.terminate_item(db, items[0].id, '取消', user.id)
    for action in [
        lambda: order_service.update_item(db, items[0].id, OrderItemUpdate(order_qty=3), user.id),
        lambda: order_service.delete_item(db, items[0].id, user.id),
        lambda: report_service._lock_order_then_item(db, items[0].id),
    ]:
        with pytest.raises(ValueError, match='已终止'):
            action()
        db.rollback()
    progress_service.recalc_item_status(db, items[0])
    assert items[0].status == C.ITEM_TERMINATED
    items[1].status = C.ITEM_SHIPPED
    progress_service.sync_order_status(db, order.id)
    assert order.status == C.ORDER_SHIPPED


def test_shipped_line_cannot_terminate(db):
    user, customer, order, items = saved_order(db)
    items[0].status = C.ITEM_SHIPPED
    db.commit()
    with pytest.raises(ValueError, match='已有发货'):
        order_service.terminate_item(db, items[0].id, '取消', user.id)
    assert order.charged_amount == 200 and customer.balance == 100


def test_later_whole_order_termination_does_not_double_refund(db):
    user, customer, order, items = saved_order(db)
    order_service.terminate_item(db, items[0].id, '取消一项', user.id)
    order_service.terminate_order(db, order.id, '取消剩余', user.id)
    assert customer.balance == 300 and order.charged_amount == 0


def test_last_active_line_cannot_be_deleted_without_refund(db):
    user, customer, order, items = saved_order(db)
    order_service.terminate_item(db, items[0].id, '取消一项', user.id)
    with pytest.raises(ValueError, match='至少保留'):
        order_service.delete_item(db, items[1].id, user.id)
    db.rollback()
    assert order.status == C.ORDER_PRODUCING
    assert order.charged_amount == 100 and customer.balance == 200


from tests.test_domestic_delivery import context, complete, ship
from tests.test_domestic_conditional_routing import conditional_route, conditional_order
from app.domestic import shipment_service
from app.domestic.models import DomesticReportLog, DomesticItemUnit, DomesticOrder


def test_work_history_and_both_scan_modes_are_preserved_but_blocked(db, context):
    ctx = context
    units, log = complete(db, ctx, 3)
    order = db.get(DomesticOrder, ctx.item.order_id)
    order_service.terminate_item(db, ctx.item.id, '取消生产', order.created_by)
    assert db.get(DomesticReportLog, log.id).revoked == 0
    assert db.query(DomesticItemUnit).filter_by(item_id=ctx.item.id, status=1).count() == ctx.item.order_qty
    for result in [report_service.scan_item(db, ctx.item.id, ctx.worker.id),
                   report_service.scan_unit(db, units[0].id, ctx.worker.id)]:
        assert result['can_submit'] is False
        assert result['block_reason'] == 'ITEM_TERMINATED'
    assert shipment_service.ready_units(db, ctx.item)[0] == []
    with pytest.raises(ValueError):
        ship(db, ctx)
    db.rollback()
    assert order.charged_amount == 0


def test_partial_delivery_blocks_termination(db, context):
    ctx = context
    complete(db, ctx, 3)
    ship(db, ctx)
    order = db.get(DomesticOrder, ctx.item.order_id)
    charge = order.charged_amount
    with pytest.raises(ValueError, match='已有发货'):
        order_service.terminate_item(db, ctx.item.id, '取消', order.created_by)
    assert order.charged_amount == charge
    assert ctx.item.status != C.ITEM_TERMINATED


def test_service_requires_current_reads_for_finance_and_state(db):
    # SQLite verifies the emitted MySQL lock contract; it cannot emulate RR snapshots.
    from sqlalchemy import event
    user, customer, order, items = saved_order(db)
    statements = []
    def capture(state):
        if state.is_select:
            statements.append(state.statement)
    event.listen(db, 'do_orm_execute', capture)
    try:
        order_service.terminate_item(db, items[0].id, '取消', user.id)
    finally:
        event.remove(db, 'do_orm_execute', capture)
    money_reads = [s for s in statements if list(s.selected_columns.keys()) == ['order_qty', 'unit_price']]
    state_reads = [s for s in statements if list(s.selected_columns.keys()) == ['status']]
    assert money_reads and state_reads
    assert all(s._for_update_arg is not None for s in money_reads + state_reads)


def test_production_order_termination_has_no_balance_movement(db):
    user, customer, order, items = saved_order(db)
    order.order_kind = 'production'
    order.order_category = order.order_type = order.order_channel = order.required_ship_date = None
    order.total_amount = order.charged_amount = 0
    db.commit()
    order_service.terminate_item(db, items[0].id, '取消备货', user.id)
    assert order.total_amount == order.charged_amount == 0
    assert customer.balance == 100
    assert db.query(DomesticCustomerLedger).count() == 0


def test_termination_endpoint_requires_write_and_creator_and_valid_reason(db):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.auth.dependencies import get_current_user
    from app.core.database import get_db
    from app.domestic.router import router
    user, customer, order, items = saved_order(db)
    current = {'sub': str(user.id), 'permissions': [], 'roles': []}
    app = FastAPI()
    app.include_router(router, prefix='/api/domestic')
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: current
    with TestClient(app) as client:
        url = f'/api/domestic/items/{items[0].id}/terminate'
        assert client.post(url, json={'reason': '取消'}).status_code == 403
        current['permissions'] = ['domestic:write']
        assert client.post(url, json={'reason': 'x' * 501}).status_code == 422
        current['sub'] = str(user.id + 1)
        assert client.post(url, json={'reason': '取消'}).status_code == 400
        current['sub'] = str(user.id)
        response = client.post(url, json={'reason': '取消'})
        assert response.status_code == 200 and response.json()['code'] == 200
        assert items[1].status == C.ITEM_PRODUCING


def test_cancelled_line_exports_as_cancelled_with_no_fulfillment_or_subtotal(db):
    from app.domestic.export_service import build_order_workbook
    from openpyxl import load_workbook
    user, customer, order, items = saved_order(db)
    order_service.terminate_item(db, items[0].id, '取消', user.id)
    detail = order_service.get_order_detail(db, order.id)
    workbook = load_workbook(build_order_workbook(detail))
    sheet = workbook.worksheets[0]
    row = next(row for row in sheet.iter_rows() if row[0].value == 'A1（已终止）')
    assert row[4].value == '已终止'
    assert row[9].value == 0
