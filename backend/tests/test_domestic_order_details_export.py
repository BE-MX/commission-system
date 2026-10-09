from datetime import date, datetime
from io import BytesIO
from urllib.parse import unquote

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.auth.dependencies import get_current_user
from app.core.database import get_db
from app.domestic import order_service, order_details_export_service as service
from app.domestic.models import DomesticOrder, DomesticOrderItem
from app.domestic.router import router
from tests.test_domestic_order_channel_source import context
from tests.test_domestic_customer_order_controls import _user, _customer


@pytest.fixture
def orders(db):
    result = []
    first_user, first_customer, original_payload = context(db, 'detail-export-0')
    created = order_service.create_order(db, original_payload, first_user.id)
    first_order = db.get(DomesticOrder, created['id'])
    first_item = db.query(DomesticOrderItem).filter_by(order_id=first_order.id).one()
    for index, name in enumerate(('张三', '李四')):
        user = first_user if index == 0 else _user(db, 'detail-export-1')
        customer = first_customer if index == 0 else _customer(db, user, 'second-shop')
        user.real_name = name
        customer.customer_source = 'referral'
        customer.owner_user_id = user.id
        if index == 0:
            order, item = first_order, first_item
        else:
            values = {c.name: getattr(first_order, c.name) for c in first_order.__table__.columns
                      if c.name not in ('id', 'domestic_no', 'request_id', 'customer_id', 'created_by')}
            order = DomesticOrder(**values, domestic_no='DO-SECOND', request_id='detail-export-1',
                                  customer_id=customer.id, created_by=user.id)
            db.add(order)
            db.flush()
            values = {c.name: getattr(first_item, c.name) for c in first_item.__table__.columns if c.name not in ('id', 'order_id')}
            item = DomesticOrderItem(**values, order_id=order.id)
            db.add(item)
        order.order_no = f'000{index}'
        order.order_date = date(2026, 9, 2 - index)
        order.order_type = 'first_order'
        order.order_channel = 'wechat'
        order.created_at = datetime(2026, 9, 2 - index, 10)
        item.attrs_snapshot = {'product_type': 'cap', 'length': '20厘米', 'density': '中',
                               'craft': '递顶', 'hair_style_series': '直发', 'size': 'L', 'net_color': '浅棕'}
        item.order_qty = index + 1
        result.append((user, customer, order, item))
    db.commit()
    return result


def workbook(db, **filters):
    return load_workbook(service.export_order_details(db, **filters))


def test_template_columns_sales_sheets_and_ascending_dates(db, orders):
    wb = workbook(db)
    assert wb.sheetnames == ['总表', '李四', '张三']
    total = wb['总表']
    assert [c.value for c in total[1]] == [
        '序号', '订单状态', '下单日期', '归属销售', '客户订单号', '客户店名',
        '头套/发片', '发长', '发量', '头套工艺/发片工艺尺寸', '发型系列', '头套尺码', '网帽颜色', '下单数量',
    ]
    assert [total.cell(r, 3).value.date() for r in (2, 3)] == [date(2026, 9, 1), date(2026, 9, 2)]
    assert [total.cell(r, 4).value for r in (2, 3)] == ['李四', '张三']
    assert [total.cell(r, 5).value for r in (2, 3)] == ['0001', '0000']
    assert [total.cell(2, c).value for c in range(7, 15)] == ['头套', '20厘米', '中', '递顶', '直发', 'L', '浅棕', 2]
    for ws in wb:
        assert ws['A2'].value == 1
        assert ws.max_row == (3 if ws.title == '总表' else 2)
        assert ws['A1'].font.name == '微软雅黑'
        assert ws['A1'].font.bold


def test_combined_filters_match_list_and_single_sales_omits_total(db, orders):
    user, customer, order, _ = orders[0]
    filters = dict(keyword='0000', status=0, customer_id=customer.id, customer_name=customer.shop_name,
                   order_kind='business', order_category='normal', order_type='first_order',
                   order_channel='wechat', customer_source='referral', owner_user_id=user.id,
                   date_start=date(2026, 9, 2), date_end=date(2026, 9, 2),
                   include_all=False, creator_id=user.id)
    listed, total = order_service.list_orders(db, **filters)
    assert total == 1 and listed[0]['id'] == order.id
    wb = workbook(db, **filters)
    assert wb.sheetnames == ['张三']
    assert wb.active['E2'].value == '0000'
    for key, value in [('keyword', 'missing'), ('status', 3), ('customer_id', 99999),
                       ('customer_name', 'missing'), ('order_kind', 'production'),
                       ('order_category', 'special'), ('order_type', 'missing'),
                       ('order_channel', 'missing'), ('customer_source', 'missing'),
                       ('owner_user_id', 99999), ('date_start', date(2026, 9, 3)),
                       ('date_end', date(2026, 9, 1)), ('creator_id', 99999)]:
        assert workbook(db, **(filters | {key: value})).active.max_row == 1


def test_all_details_beyond_page_size_and_deleted_orders_excluded(db, orders):
    user, _, order, item = orders[0]
    for index in range(2, 27):
        values = {c.name: getattr(item, c.name) for c in item.__table__.columns if c.name not in ('id', 'line_no')}
        db.add(DomesticOrderItem(**values, line_no=index))
    orders[1][2].deleted_flag = 1
    db.commit()
    wb = workbook(db, include_all=False, creator_id=user.id)
    assert wb.sheetnames == ['张三']
    assert wb.active.max_row == 27
    assert wb.active['A27'].value == 26


def test_sheet_names_distinct_and_user_text_stays_literal(db, orders):
    for user, customer, order, item in orders:
        user.real_name = '总表'
        customer.shop_name = f'=HYPERLINK("https://example.com/{customer.id}")'
        order.order_no = '+SUM(1,2)'
        item.attrs_snapshot = {'product_type': 'piece', 'craft': '8×10', 'length': '15厘米'}
    db.commit()
    wb = workbook(db)
    assert len(set(wb.sheetnames)) == 3 and wb.sheetnames[0] == '总表'
    assert wb.active['F2'].value.startswith('=HYPERLINK')
    assert wb.active['F2'].data_type == 's'
    assert wb.active['E2'].data_type == 's'
    assert [wb.active.cell(2, c).value for c in (7, 9, 10, 11, 12, 13)] == ['发片', None, '8×10', None, None, None]
    assert service.sheet_title('[]:/\\?*' + '名' * 40, set()).startswith('_______')
    assert len(service.sheet_title('名' * 40, set())) == 31


def test_production_without_owner_and_empty_export(db, orders):
    order = orders[0][2]
    order.order_kind = 'production'
    order.customer_id = None
    order.order_category = order.order_type = order.order_channel = order.required_ship_date = None
    order.total_amount = order.charged_amount = 0
    db.commit()
    wb = workbook(db, order_kind='production')
    assert wb.sheetnames == ['未归属销售']
    assert wb.active['D2'].value == '未归属销售'
    empty = workbook(db, keyword='no-match')
    assert empty.sheetnames == ['总表'] and empty.active.max_row == 1


def test_illegal_control_characters_do_not_break_batch_export(db, orders):
    user, customer, order, item = orders[1]
    user.real_name = '李\x0b四'
    customer.shop_name = '门\x0c店\n名称'
    order.order_no = '000\x0101'
    item.attrs_snapshot = {'product_type': 'cap', 'craft': '工\x0b艺\t名称'}
    db.commit()
    wb = workbook(db)
    assert wb.active['D2'].value == '李四'
    assert wb.active['E2'].value == '00001'
    assert wb.active['F2'].value == '门店\n名称'
    assert wb.active['J2'].value == '工艺\t名称'


def test_sales_names_matching_copy_names_keep_their_exact_names(db, orders):
    orders[1][0].real_name = '归属销售 Copy1'
    orders[0][0].real_name = '总表 Copy'
    db.commit()
    assert workbook(db).sheetnames == ['总表', '归属销售 Copy1', '总表 Copy']


def test_same_date_orders_use_creation_time_then_stable_item_sequence(db, orders):
    orders[0][2].order_date = orders[1][2].order_date
    orders[0][2].created_at = datetime(2026, 9, 1, 8)
    orders[1][2].created_at = datetime(2026, 9, 1, 10)
    db.commit()
    assert [workbook(db).active.cell(r, 5).value for r in (2, 3)] == ['0000', '0001']


def test_endpoint_authorization_visibility_filters_and_download(db, orders):
    user = orders[0][0]
    current = {'sub': str(user.id), 'roles': [], 'permissions': ['domestic:read']}
    app = FastAPI()
    app.include_router(router, prefix='/api/domestic')
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: current
    with TestClient(app) as client:
        response = client.get('/api/domestic/orders/export-details', params={'page': 2, 'page_size': 1, 'sort_order': 'descending'})
        assert response.status_code == 200
        assert workbook_bytes(response).sheetnames == ['张三']
        assert '内贸订单明细-' in unquote(response.headers['content-disposition'])
        current['permissions'] = ['domestic:read', 'domestic:read_all']
        assert workbook_bytes(client.get('/api/domestic/orders/export-details')).sheetnames == ['总表', '李四', '张三']
        filtered = client.get('/api/domestic/orders/export-details', params={'owner_user_id': user.id, 'date_end': '2026-09-02'})
        assert workbook_bytes(filtered).sheetnames == ['张三']
        assert client.get('/api/domestic/orders/export-details', params={'owner_user_id': 0}).status_code == 422
        current['permissions'] = []
        assert client.get('/api/domestic/orders/export-details').status_code == 403


def workbook_bytes(response):
    assert response.status_code == 200
    return load_workbook(BytesIO(response.content))
