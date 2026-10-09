"""按列表筛选导出模板明细；查询不分页、不读取或修复生产进度。"""

from collections import OrderedDict
from copy import copy
from io import BytesIO
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from sqlalchemy.orm import Session

from app.auth.models import ArkUser
from app.domestic import constants as C
from app.domestic.models import DomesticCustomer, DomesticOrder, DomesticOrderItem
from app.domestic.order_query_service import filtered_order_query


_TEMPLATE = Path(__file__).resolve().parents[2] / 'assets/domestic/order_details_template.xlsx'


def sheet_title(name: str, used: set[str]) -> str:
    """Excel names are case insensitive, limited to 31 chars and exclude punctuation."""
    base = re.sub(r'[\[\]:/\\?*\x00-\x1f]', '_', name).strip().strip("'") or '未归属销售'
    title = base[:31].rstrip("'")
    suffix = 1
    while title.casefold() in used or title.casefold() == 'history':
        suffix += 1
        tail = f'（{suffix}）'
        title = base[:31 - len(tail)] + tail
    used.add(title.casefold())
    return title


def export_order_details(db: Session, **filters) -> BytesIO:
    query = filtered_order_query(db, **filters)
    rows = query.with_entities(
        DomesticOrder.status, DomesticOrder.order_date, DomesticOrder.order_no,
        DomesticCustomer.owner_user_id, ArkUser.real_name.label('owner_name'),
        DomesticCustomer.shop_name, DomesticOrderItem.attrs_snapshot, DomesticOrderItem.order_qty,
    ).join(DomesticOrderItem, DomesticOrderItem.order_id == DomesticOrder.id).outerjoin(
        DomesticCustomer, DomesticCustomer.id == DomesticOrder.customer_id,
    ).outerjoin(ArkUser, ArkUser.id == DomesticCustomer.owner_user_id).order_by(
        DomesticOrder.order_date.asc(), DomesticOrder.created_at.asc(), DomesticOrder.id.asc(),
        DomesticOrderItem.line_no.asc(), DomesticOrderItem.id.asc(),
    ).all()
    groups = OrderedDict()
    values = []
    for row in rows:
        owner_name = row.owner_name or (f'未命名销售-{row.owner_user_id}' if row.owner_user_id else '未归属销售')
        attrs = row.attrs_snapshot or {}
        cap = attrs.get('product_type') == 'cap'
        data = (
            C.ORDER_STATUS_LABELS.get(row.status, str(row.status)), row.order_date, owner_name,
            row.order_no, row.shop_name, C.PRODUCT_TYPES.get(attrs.get('product_type'), attrs.get('product_type')),
            attrs.get('length'), attrs.get('density') if cap else None, attrs.get('craft'),
            attrs.get('hair_style_series') if cap else None, attrs.get('size') if cap else None,
            attrs.get('net_color') if cap else None, row.order_qty,
        )
        values.append(data)
        groups.setdefault(row.owner_user_id, (owner_name, []))[1].append(data)

    wb = load_workbook(_TEMPLATE)
    sales_template, total_template = wb['归属销售'], wb['总表']
    used = {'总表'.casefold()}
    sheets = []

    def add_sheet(template, title, data):
        ws = wb.copy_worksheet(template)
        sheets.append((ws, title))
        styles = [copy(cell._style) for cell in template[2]]
        ws.delete_rows(2, ws.max_row - 1)
        for index, values_row in enumerate(data, start=1):
            for col, value in enumerate((index, *values_row), start=1):
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub('', value)
                cell = ws.cell(index + 1, col, value)
                cell._style = copy(styles[col - 1])
                # Keep customer IDs and user supplied strings literal, including leading '='.
                if isinstance(value, str):
                    cell.data_type = 's'
                if col == 3:
                    cell.number_format = 'yyyy/mm/dd'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:N{ws.max_row}'
        ws.print_title_rows = '1:1'
        ws.print_area = f'A1:N{ws.max_row}'
        return ws

    if len(groups) != 1:
        add_sheet(total_template, '总表', values)
    for name, data in groups.values():
        add_sheet(sales_template, sheet_title(name, used), data)
    wb.remove(sales_template)
    wb.remove(total_template)
    # Avoid collisions with other copies' provisional names during final renaming.
    temporary_names = {ws.title.casefold() for ws in wb} | {title.casefold() for _, title in sheets}
    for index, (ws, _) in enumerate(sheets):
        ws.title = sheet_title(f'_export_tmp_{index}', temporary_names)
    for ws, title in sheets:
        ws.title = title
    wb.active = 0
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
