"""Employee PI exports using the supplied LeShine workbook, without repricing."""
from copy import copy
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile
import math
import unicodedata

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from app.invoice.export_service import _excel_value
from app.invoice.presale_lines import editor_document

TEMPLATE_PATH = Path(__file__).resolve().parent / "assets" / "portal_pi_template.xlsx"


def document(invoice):
    """Only invoice display fields; never internal receipts, balances or price rules."""
    source = editor_document(invoice)
    currency = source.currency or ""
    items = []
    for item in source.items:
        description = item.product_display or item.product_name or ""
        details = [str(v) for v in (item.model, item.curl) if v]
        if item.discount_amount:
            details.append(f"Discount: {item.discount_amount:.2f} {currency}")
        if details:
            description += "\n" + " / ".join(details)
        items.append([len(items) + 1, description, item.length or "", item.color or "",
            item.net_weight_grams or "", item.quantity, item.price_per_piece, item.total_price])
    fees = [("Products Cost" if any(i.product_kind == "accessory" for i in source.items) else "Hair Cost", source.product_amount),
        ("Shipping Cost", source.shipping_fee), (source.surcharge_name or "Payment Surcharge", source.surcharge_amount)]
    if source.internal_accessory or source.packaging_quantity:
        fees.append((f"Packaging ({source.packaging_quantity or 0})", source.internal_accessory or Decimal("0")))
    fees.append((f"Total ({currency})", source.total_amount))
    recipient = " / ".join(dict.fromkeys(str(v) for v in (source.customer_name, source.contact_name) if v))
    return dict(invoice_no=source.invoice_no, invoice_date=source.invoice_date, currency=currency,
        recipient=recipient, phone=source.contact_phone or "", email=source.contact_email or "",
        address=source.delivery_address or "", salesperson=source.sales_user_name or "",
        sales_phone=source.sales_phone or "", sales_email=source.sales_email or "", items=items, fees=fees,
        payment=source.payment_term or "", shipping=source.express_channel or "", notes=source.remark or "")


def _line_count(value, width):
    return sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line) / width))
        for line in str(value or "").split("\n"))


def _sheet(data):
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.worksheets[0]
    for extra in wb.worksheets[1:]:
        wb.remove(extra)
    ws.title = "PI"
    # openpyxl retains cell styles/images, but discards floating text boxes. They
    # are restored from the source ZIP by build_workbook after saving.
    styles = {row: [copy(ws.cell(row, col)._style) for col in range(1, 9)] for row in (12, 17, 20, 21)}
    for row in ws:
        for cell in row:
            if cell.hyperlink:
                cell.hyperlink = None  # Do not retain the example salesperson's mailto.
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= 12:
            ws.unmerge_cells(str(merged))
    ws.delete_rows(12, ws.max_row)
    for row in list(ws.row_dimensions):
        if row >= 12:
            del ws.row_dimensions[row]
    headers = {"C6": data["recipient"], "C7": data["phone"], "C8": data["email"], "C9": data["address"],
        "G6": data["salesperson"], "G7": data["sales_phone"], "G8": data["sales_email"],
        "G9": data["invoice_date"], "G10": data["invoice_no"]}
    for address, value in headers.items():
        cell = ws[address]
        cell.value = _excel_value(value)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["G9"].number_format = "yyyy/m/d"
    for row in (6, 7, 8):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 20,
            14 * max(_line_count(ws.cell(row, 3).value, 30), _line_count(ws.cell(row, 7).value, 18)))
    address_height = max(71, _line_count(data["address"], 30) * 13)
    ws.row_dimensions[9].height = address_height / 2
    ws.row_dimensions[10].height = max(address_height / 2, _line_count(data["invoice_no"], 18) * 14)
    for index in range(max(5, len(data["items"]))):
        row = 12 + index
        values = data["items"][index] if index < len(data["items"]) else [None] * 8
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col)
            cell._style = copy(styles[12][col - 1])
            cell.value = float(value) if isinstance(value, Decimal) else _excel_value(value)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if col in (7, 8):
                cell.number_format = "0.00##" if col == 7 else "0.00"
        ws.row_dimensions[row].height = max(29, 12 * max(_line_count(values[col], width)
            for col, width in ((1, 32), (2, 8), (3, 19), (4, 6))))
    row = 12 + max(5, len(data["items"]))
    for label, amount in data["fees"]:
        is_total = label.startswith("Total (")
        for col in range(1, 9):
            ws.cell(row, col)._style = copy(styles[20 if is_total else 17][col - 1])
        ws.cell(row, 2, _excel_value(label))
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
        ws.cell(row, 8, float(amount) if amount is not None else None).number_format = "0.00"
        ws.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = max(23, _line_count(label, 23) * 20)
        row += 1
    for label, value in (("Payment term", data["payment"]), ("Shipping method", data["shipping"]), ("Notes", data["notes"])):
        if not value and label != "Payment term":
            continue
        for col in range(1, 9):
            ws.cell(row, col)._style = copy(styles[21][col - 1])
        ws.cell(row, 2, label)
        ws.cell(row, 3, _excel_value(value))
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        ws.cell(row, 3).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = max(26, _line_count(value, 65) * 19)
        row += 1
    ws.print_area = f"A1:H{row - 1}"
    ws.print_title_rows = "1:11"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return wb


def build_workbook(invoice):
    generated = BytesIO()
    _sheet(document(invoice)).save(generated)
    generated.seek(0)
    output = BytesIO()
    with ZipFile(TEMPLATE_PATH) as source, ZipFile(generated) as cells, ZipFile(output, "w", ZIP_DEFLATED) as result:
        artwork = {name: source.read(name) for name in source.namelist()
            if not name.endswith("/") and name.startswith(("xl/drawings/", "xl/media/"))}
        for entry in cells.infolist():
            result.writestr(entry, artwork.pop(entry.filename, cells.read(entry.filename)))
        for name, content in artwork.items():
            result.writestr(name, content)
    output.seek(0)
    return output


def build_pdf(invoice):
    # Explicit portable rendering of this template; production requires no Excel
    # or LibreOffice process. Both formats share the same stored-value mapping.
    from app.invoice.pi_template_pdf import render
    return render(document(invoice))


def artwork():
    """Text boxes and image from the supplied template, not retyped brand data."""
    ns = {"x": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with ZipFile(TEMPLATE_PATH) as source:
        root = ElementTree.fromstring(source.read("xl/drawings/drawing1.xml"))
        texts = []
        for shape in root.findall(".//x:sp", ns):
            paragraphs = tuple("".join(t.text or "" for t in p.findall(".//a:t", ns)).strip()
                for p in shape.findall(".//a:p", ns))
            if paragraphs not in texts:
                texts.append(paragraphs)
        return source.read("xl/media/image1.png"), texts
