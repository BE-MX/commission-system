"""A4 rendering of the supplied PI's eight-column grid and letterhead."""
from html import escape
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import get_settings
from app.portal.pi_pdf import registered_font
from app.invoice.pi_template import TEMPLATE_PATH, artwork


def render(data):
    font = registered_font(str(get_settings().PDF_CJK_FONT_PATH))
    # Use the supplied column proportions, including the narrow numbered column.
    template = load_workbook(TEMPLATE_PATH)
    widths = [(template.active.column_dimensions[get_column_letter(c)].width * 7 + 5) * .75 for c in range(1, 9)]
    factor = (A4[0] - 28) / sum(widths)
    widths = [w * factor for w in widths]
    body = ParagraphStyle("pi-cell", fontName=font, fontSize=9 * factor, leading=12 * factor,
        alignment=TA_CENTER, wordWrap="CJK", textColor=colors.black)
    left = ParagraphStyle("pi-left", parent=body, alignment=TA_LEFT)
    label = ParagraphStyle("pi-label", parent=body, fontSize=11 * factor, leading=14 * factor)

    def p(value, style=body):
        return Paragraph(escape(str(value if value is not None else "")).replace("\n", "<br/>"), style)

    def number(value, unit=False):
        if value is None:
            return ""
        value = f"{value:.4f}".rstrip("0").rstrip(".") if unit else f"{value:.2f}"
        if unit and "." not in value:
            value += ".00"
        return value

    metadata = [
        ["", "To:", data["recipient"], "", "", "From:", data["salesperson"], ""],
        ["", "TEL/Fax:", data["phone"], "", "", "TEL:", data["sales_phone"], ""],
        ["", "E-mail:", data["email"], "", "", "E-mail:", data["sales_email"], ""],
        ["", "Delivery\naddress:", data["address"], "", "", "Date:", str(data["invoice_date"] or ""), ""],
        ["", "", "", "", "", "No.:", data["invoice_no"], ""],
    ]
    rows = [[p(value, label if c in (1, 5) else body) for c, value in enumerate(row)] for row in metadata]
    rows.append([p(v) for v in ("", "Product", "Length", "Color", "Weight", "Quantity", "Unit Price", "Total Price")])
    for item in data["items"]:
        values = [*item[:6], number(item[6], True), number(item[7])]
        rows.append([p(v, left if c == 1 else body) for c, v in enumerate(values)])
    for _ in range(max(0, 5 - len(data["items"]))):
        rows.append([""] * 8)
    commands = [("GRID", (1, 0), (-1, -1), .5, colors.black),
        ("GRID", (0, 5), (-1, -1), .5, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("SPAN", (1, 3), (1, 4)), ("SPAN", (2, 3), (4, 4))]
    for row in range(5):
        if row < 3:
            commands.append(("SPAN", (2, row), (4, row)))
        commands.append(("SPAN", (6, row), (7, row)))
    table = Table(rows, colWidths=widths, repeatRows=6, splitByRow=1, splitInRow=1, hAlign="LEFT")
    table.setStyle(TableStyle(commands))
    summaries = []
    for title, amount in data["fees"]:
        summaries.append(["", p(title, label), "", "", "", "", "", p(number(amount))])
    for title, value in (("Payment term", data["payment"]), ("Shipping method", data["shipping"])):
        if value or title == "Payment term":
            summaries.append(["", p(title, label), p(value, left), "", "", "", "", ""])
    summary = Table(summaries, colWidths=widths, hAlign="LEFT")
    summary_commands = [("GRID", (1, 0), (-1, -1), .5, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]
    for row in range(len(summaries)):
        summary_commands.append(("SPAN", (2, row), (6 if row < len(data["fees"]) else 7, row)))
    summary.setStyle(TableStyle(summary_commands))
    logo, textboxes = artwork()

    def page(canvas, doc):
        canvas.saveState()
        top = A4[1] - 22
        # Header image is the original embedded logo, with its alpha channel.
        canvas.drawImage(ImageReader(BytesIO(logo)), 14 + widths[0] + 12 * factor, top - 55 * factor,
            width=120 * factor, height=48 * factor, mask="auto")
        x = 14 + widths[0] + widths[1]
        company = next((t for t in textboxes if any("LESHINE" in line for line in t)), ())
        y = top
        for index, line in enumerate(company):
            size = (11 if index < 2 else 9) * factor
            style = ParagraphStyle("letterhead", parent=left, fontSize=size, leading=size + 3)
            para = p(line, style)
            _, height = para.wrap(sum(widths[2:6]), 200)
            para.drawOn(canvas, x, y - height)
            y -= height + 1
        title = next((t for t in textboxes if any("Proforma" in line for line in t)), ("Proforma", "Invoice"))
        y = top - 4
        for line in title:
            para = p(line, ParagraphStyle("pi-title", parent=body, fontSize=13 * factor, leading=17 * factor))
            _, height = para.wrap(sum(widths[6:]), 100)
            para.drawOn(canvas, 14 + sum(widths[:6]), y - height)
            y -= height
        canvas.setFont(font, 8)
        canvas.drawRightString(A4[0] - 14, 16, f"{data['invoice_no']}  /  {doc.page}")
        canvas.restoreState()

    output = BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=A4, leftMargin=14, rightMargin=14, topMargin=112,
        bottomMargin=32, title=f"{data['invoice_no']} - Proforma Invoice", author="LeShine")
    story = [table, KeepTogether(summary)]
    if data["notes"]:
        story += [Spacer(1, 10), p("Notes", label), p(data["notes"], left)]
    pdf.build(story, onFirstPage=page, onLaterPages=page)
    output.seek(0)
    return output
