"""Editable DOCX outbound sheet matching the A4 print template's physical dimensions."""
import io
import re

import qrcode
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.core.time import beijing_now
from app.shipping_inspection.print_service import sort_outbound_print_items

WIDTHS = [7.92, 35.64, 64.35, 45.54, 13.86, 30.69]  # 198mm × 4/18/32.5/23/7/15.5%


def _xml(tag, **attrs):
    element = OxmlElement(f"w:{tag}")
    for key, value in attrs.items():
        element.set(qn(f"w:{key}"), str(value))
    return element


def _text(value):
    # XML 1.0 disallows C0 controls except tab/newline/CR.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))


def _run(paragraph, value, size=9.75, bold=False):
    run = paragraph.add_run(_text(value))
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    return run


def _cell(cell, value, size=9.75, bold=False, shade=None, center=False, spec=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if spec:
        for part in re.split(r"((?<![A-Za-z0-9])B[13](?![A-Za-z0-9]))", _text(value), flags=re.I):
            emphasized = bool(re.fullmatch(r"B[13]", part, re.I))
            _run(paragraph, part, 12 if emphasized else size, emphasized)
    else:
        _run(paragraph, value, size, bold)
    if shade:
        cell._tc.get_or_add_tcPr().insert_element_before(
            _xml("shd", val="clear", fill=shade), "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _table(parent, widths, rows=1, borders=True):
    table = parent.add_table(rows=rows, cols=len(widths))
    table.autofit = False
    if borders:
        table.style = "Table Grid"
    props = table._tbl.tblPr
    props.find(qn("w:tblW")).set(qn("w:w"), str(Mm(sum(widths)).twips))
    props.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    margins = _xml("tblCellMar")
    for side, points in {"top": 4.5, "left": 3.75, "bottom": 4.5, "right": 3.75}.items():
        margins.append(_xml(side, w=int(points * 20), type="dxa"))
    props.insert_element_before(margins, "w:tblLook", "w:tblCaption", "w:tblDescription")
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Mm(width)
    return table


def _line(paragraph):
    borders = _xml("pBdr")
    borders.append(_xml("bottom", val="single", sz=8, color="000000", space=5))
    paragraph._p.get_or_add_pPr().insert_element_before(
        borders, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:contextualSpacing", "w:jc", "w:rPr")


def build_outbound_word(record: dict, items: list[dict], qr_data: str) -> bytes:
    doc = Document()
    doc.settings.element.find(qn("w:zoom")).set(qn("w:percent"), "100")
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(6)
    section.top_margin = section.bottom_margin = Mm(12)
    normal = doc.styles["Normal"]
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    normal.font.name, normal.font.size = "Microsoft YaHei", Pt(9.75)
    title = doc.add_paragraph()
    _run(title, "出库单", 16.5, True)
    number = doc.add_paragraph()
    _run(number, f"单号：{record.get('outbound_no', '')}")
    _line(number)
    number.paragraph_format.space_after = Pt(10)

    head = _table(doc, [24, 137, 37], rows=3)
    for row, (label, key) in zip(head.rows, [("客户名称", "customer_name"), ("出库日期", "outbound_date"), ("负责人", "owner_name")]):
        _cell(row.cells[0], label, 10.5, shade="F0F0F0")
        _cell(row.cells[1], record.get(key), 10.5, bold=key == "customer_name")
    qr_cell = head.cell(0, 2).merge(head.cell(2, 2))
    qr_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    image = io.BytesIO()
    qrcode.make(qr_data).save(image, format="PNG")
    image.seek(0)
    qr_cell.paragraphs[0].add_run().add_picture(image, width=Mm(31.75))
    hint = qr_cell.add_paragraph()
    hint.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(hint, "扫码查看出库信息", 8.25)
    # QR block is alongside the metadata, without a surrounding frame in the print template.
    borders = _xml("tcBorders")
    for side in ("top", "left", "bottom", "right"):
        borders.append(_xml(side, val="nil"))
    qr_cell._tc.get_or_add_tcPr().insert_element_before(borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    remark = _table(doc, [198])
    _cell(remark.cell(0, 0), "发货备注", bold=True)
    _run(remark.cell(0, 0).add_paragraph(), record.get("remark") or "无")
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(9)
    heading.paragraph_format.space_after = Pt(4.5)
    _run(heading, "出库明细", 10.5, True)
    if items:
        table = _table(doc, WIDTHS, rows=len(items) + 1)
        table.rows[0]._tr.get_or_add_trPr().append(_xml("tblHeader"))
        for index, label in enumerate(["#", "产品类别", "规格", "颜色/尺寸/克重", "数量", "批次号"]):
            _cell(table.cell(0, index), label, bold=True, shade="F0F0F0", center=index == 4)
        for index, item in enumerate(sort_outbound_print_items(items), start=1):
            parts = re.split(r"[/／]", _text(item.get("product_name")).strip(), maxsplit=1)
            values = [index, parts[0].strip(), item.get("spec"), parts[1].strip() if len(parts) > 1 else "", str(item.get("qty", "")), ""]
            row = table.rows[index]
            row.height, row.height_rule = Mm(12), WD_ROW_HEIGHT_RULE.AT_LEAST
            row._tr.get_or_add_trPr().append(_xml("cantSplit"))
            for column, value in enumerate(values):
                _cell(row.cells[column], value, size=9 if column == 1 else 10.5 if column == 3 else 9.75,
                      bold=column == 3, shade="F5F5F5" if index % 2 == 0 else None, center=column == 4, spec=column == 2)
    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(9)
    _line(footer)
    _run(footer, f"莱莎方舟平台 · 发货检验    打印时间：{beijing_now():%Y-%m-%d %H:%M:%S}    单号：{record.get('outbound_no', '')}", 8.25).font.color.rgb = RGBColor.from_string("555555")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
