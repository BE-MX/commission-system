"""薪资模块 — 导出（M4）：双 sheet 工资表 + 银行代发盘 + 导出留痕与作废判定。

设计文档 §7.3 / §8.1 M4。三条不能松的规矩：

1. **锁定批次只读快照**。confirmed 批次的姓名/工号/部门/职务/入职日期/卡号/参保主体
   一律取 `SalaryRecord.snap_*`，这里**不 import、不查**活档案。锁定后有人改了档案
   （换卡、调岗），已锁定的月份导出来必须还是锁定那一刻的样子——否则就是 §2.5
   错误 2（卡号张冠李戴）的系统版。未冻结的锁定批次（早于快照机制）直接拒导，
   不回落活档案。未锁定批次可以导「预览稿」，读活档案，文件上明写不得作为发放依据。

2. **每次导出先留痕，留痕失败就不给文件**。作废判定（决策 A4）靠的是导出记录：
   没留下记录的文件，事后既查不到也作废不了。所以导出事件与文件在同一个事务里
   生成：批次行发号 → 写事件 → 生成文件（编号印进文件）→ commit 成功才返回文件。
   不走 `period_service.log_event`——那个是「留痕失败不连累主流程」的语义，这里相反。

3. **明文卡号不进日志、不进事件 payload**。明文只在生成文件的那一刻解密出来，
   payload 只记「是否含明文」这个布尔；权限判断在路由层（salary:admin）。

作废判定：每条锁定状态下的导出记下当时批次的 `status_version`。离开 confirmed 的
唯一出口是 unlock（跃迁白名单 `_ALLOWED`），所以「批次现在的版本比导出时高」就等于
「导出之后解锁过」——不靠 unlock 事件（它走 log_event，可能缺）、不靠时钟（秒级、
提交顺序不定）。unlock 事件只用来给出作废的时间与原因。

并发：导出事务先对批次行做 `export_seq + 1` 的 UPDATE——拿到行锁、编号在批次内唯一，
并与解锁/锁定的条件 UPDATE 串行；拿锁后若发现状态或版本已不是取数时的样子（取数与
拿锁之间有人解锁），整笔回滚让用户重试，不留一份「解锁前读的数、却记成有效」的导出。

已经发出去的文件改不了，作废落在三处：导出历史标「已作废」并给出解锁原因；新文件
的脚注写明「替代第 N 次导出」；每份文件都印着导出编号，纸面件可以回系统核对。
"""

from __future__ import annotations

import csv
import hashlib
import io
import logging
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Optional
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select, update
from sqlalchemy.orm import Session

from app.core.time import beijing_now
from app.salary import period_service, service
from app.salary.models import (
    SalaryDeptMapping,
    SalaryEmployeeProfile,
    SalaryPeriod,
    SalaryPeriodEvent,
    SalaryRecord,
)
from app.salary.pii import decrypt_pii, mask_pii, normalize_bank_card

logger = logging.getLogger("commission")

KIND_WORKBOOK = "workbook"
KIND_BANK = "bank"
KIND_LABELS = {KIND_WORKBOOK: "工资表", KIND_BANK: "银行代发盘"}

EVENT_EXPORT = "export"
NO_GROUP = "未分组"
NO_ENTITY = "未登记参保主体"

# 能导出工资表的状态：算过才有明细。预览稿（未锁定）同样要求算过。
_EXPORTABLE = {
    period_service.STATUS_CALCULATED,
    period_service.STATUS_REVIEWING,
    period_service.STATUS_CONFIRMED,
}

EXPORT_STATUS_LABELS = {
    "valid": "有效",
    "voided": "已作废",
    "preview": "预览稿",
}


class SalaryExportError(ValueError):
    """导出前置条件不满足（状态不对、快照缺失、银行盘有不能发的人）。文案直接给用户看。"""


@dataclass
class ExportRow:
    """一行导出数据。来源（快照 / 活档案）在 load_export_rows 一处决定，下游不再分支。"""

    employee_id: int
    emp_no: Optional[str]
    name: Optional[str]
    dept_detail: Optional[str]
    dept_group: Optional[str]
    position: Optional[str]
    hire_date: Optional[date]
    bank_card_cipher: Optional[str]
    insurance_entity: Optional[str]
    record: SalaryRecord


# ---------------------------------------------------------------------------
# 取数
# ---------------------------------------------------------------------------

def is_locked(period: SalaryPeriod) -> bool:
    return period.status == period_service.STATUS_CONFIRMED


def load_export_rows(db: Session, period: SalaryPeriod) -> list[ExportRow]:
    """按明细序号取整批导出行。锁定批次只读快照列，见模块注释第 1 条。"""
    if period.status not in _EXPORTABLE:
        raise SalaryExportError("批次还没有计算，没有可导出的工资明细；请先完成计算")
    records = (
        db.query(SalaryRecord)
        .filter(SalaryRecord.period_id == period.id)
        .order_by(SalaryRecord.seq_no, SalaryRecord.id)
        .all()
    )
    if not records:
        raise SalaryExportError("批次里没有工资明细，请先计算")

    if is_locked(period):
        unfrozen = [r for r in records if r.snapshot_at is None]
        if unfrozen:
            # 不回落活档案：那正是快照要防的事。解锁重锁会按当前档案补冻。
            raise SalaryExportError(
                f"该批次有 {len(unfrozen)} 行明细在锁定时没有冻结档案快照（锁定早于快照机制），"
                "无法按锁定时的信息导出。请管理员解锁后重新锁定，再导出"
            )
        return [
            ExportRow(
                employee_id=r.employee_id,
                emp_no=r.snap_emp_no,
                name=r.snap_name,
                dept_detail=r.snap_dept_detail,
                dept_group=r.snap_dept_group,
                position=r.snap_position,
                hire_date=r.snap_hire_date,
                bank_card_cipher=r.snap_bank_card_cipher,
                insurance_entity=r.snap_insurance_entity,
                record=r,
            )
            for r in records
        ]

    profiles = {
        p.id: p for p in db.query(SalaryEmployeeProfile)
        .filter(SalaryEmployeeProfile.id.in_({r.employee_id for r in records}))
    }
    dept_map = service.load_dept_group_map(db)
    rows = []
    for r in records:
        p = profiles.get(r.employee_id)
        rows.append(ExportRow(
            employee_id=r.employee_id,
            emp_no=p.emp_no if p else None,
            name=p.name if p else None,
            dept_detail=p.dept_detail if p else None,
            dept_group=service.resolve_dept_group(p, dept_map) if p else None,
            position=p.position if p else None,
            hire_date=p.hire_date if p else None,
            bank_card_cipher=p.bank_card_cipher if p else None,
            insurance_entity=p.insurance_entity if p else None,
            record=r,
        ))
    return rows


def group_rows(db: Session, rows: list[ExportRow]) -> list[tuple[str, list[ExportRow]]]:
    """汇总表分组：大部门按映射表 sort_order 排，组内保持明细顺序。

    映射表只用来排组的先后（配置，不是个人信息），组归属仍取行上的 dept_group
    （锁定批次即快照）。映射表里没有的组排在后面，「未分组」垫底——它本该在异常
    面板里被处理掉，导出时单列一组让人一眼看见，而不是悄悄并进别的部门。
    """
    order: dict[str, int] = {}
    for m in db.query(SalaryDeptMapping).all():
        cur = order.get(m.dept_group)
        order[m.dept_group] = m.sort_order if cur is None else min(cur, m.sort_order)

    buckets: dict[str, list[ExportRow]] = {}
    for row in rows:
        buckets.setdefault((row.dept_group or "").strip() or NO_GROUP, []).append(row)

    def key(name: str) -> tuple[int, int, str]:
        if name == NO_GROUP:
            return (2, 0, name)
        if name in order:
            return (0, order[name], name)
        return (1, 0, name)

    return [(name, buckets[name]) for name in sorted(buckets, key=key)]


def _plain_card(row: ExportRow) -> str:
    return normalize_bank_card(decrypt_pii(row.bank_card_cipher)) if row.bank_card_cipher else ""


CARD_MISSING = "（无卡号，请核对档案）"


def _card_text(row: ExportRow, plain_card: bool) -> str:
    # 缺卡/解密失败不能留空白：人工发放表上一格空白很容易被当成「漏抄」随手补上
    card = _plain_card(row)
    if not card:
        return CARD_MISSING
    return card if plain_card else mask_pii(card, keep_head=4, keep_tail=4)


def _missing_card_names(rows: list[ExportRow]) -> list[str]:
    return [r.name or f"员工#{r.employee_id}" for r in rows if not _plain_card(r)]


def _dec(v: Any) -> Decimal:
    return Decimal(v) if v is not None else Decimal("0")


# ---------------------------------------------------------------------------
# 导出历史与作废判定
# ---------------------------------------------------------------------------

def _export_events(db: Session, period_id: int) -> list[SalaryPeriodEvent]:
    return (
        db.query(SalaryPeriodEvent)
        .filter(SalaryPeriodEvent.period_id == period_id,
                SalaryPeriodEvent.event_type == EVENT_EXPORT)
        .order_by(SalaryPeriodEvent.id)
        .all()
    )


def _unlock_events(db: Session, period_id: int) -> list[SalaryPeriodEvent]:
    return (
        db.query(SalaryPeriodEvent)
        .filter(SalaryPeriodEvent.period_id == period_id,
                SalaryPeriodEvent.event_type == "unlock")
        .order_by(SalaryPeriodEvent.status_version, SalaryPeriodEvent.id)
        .all()
    )


def _voided_by(
    export: SalaryPeriodEvent, unlocks: list[SalaryPeriodEvent],
) -> Optional[SalaryPeriodEvent]:
    """作废了这次导出的那次解锁（最早一次版本更高的 unlock），没有则 None。"""
    ver = export.status_version or 0
    for u in unlocks:
        if (u.status_version or 0) > ver:
            return u
    return None


@dataclass
class _Void:
    at: Optional[datetime]
    reason: Optional[str]


def _export_status(
    export: SalaryPeriodEvent, unlocks: list[SalaryPeriodEvent], period: SalaryPeriod,
) -> tuple[str, Optional[_Void]]:
    """模块注释「作废判定」：按版本判，unlock 事件只提供时间与原因。"""
    if not (export.payload or {}).get("locked"):
        return "preview", None
    if (period.status_version or 0) <= (export.status_version or 0):
        return "valid", None
    voider = _voided_by(export, unlocks)
    if voider:
        return "voided", _Void(voider.created_at, voider.reason)
    # unlock 留痕丢了：作废是确定的，原因如实说缺，不拿最近一次解锁的原因顶替
    return "voided", _Void(None, "解锁留痕缺失，原因未记录")


def export_history(db: Session, period: SalaryPeriod) -> list[dict[str, Any]]:
    """导出历史（新到旧），每条带有效 / 已作废 / 预览稿状态与作废原因。"""
    events = _export_events(db, period.id)
    unlocks = _unlock_events(db, period.id)
    names = period_service.operator_names(db, events)
    out = []
    for e in reversed(events):
        p = e.payload or {}
        status, voider = _export_status(e, unlocks, period)
        out.append({
            "id": e.id,
            "export_no": p.get("export_no"),
            "kind": p.get("kind"),
            "kind_label": KIND_LABELS.get(p.get("kind"), p.get("kind")),
            "entity": p.get("entity"),
            "plain_card": bool(p.get("plain_card")),
            "locked": bool(p.get("locked")),
            "status": status,
            "status_label": EXPORT_STATUS_LABELS[status],
            "voided_at": voider.at.isoformat() if voider and voider.at else None,
            "void_reason": voider.reason if voider else None,
            "record_count": p.get("record_count"),
            "total_amount": p.get("total_amount"),
            "filename": p.get("filename"),
            "sha256": p.get("sha256"),
            "operator_name": names.get(e.created_by) or (str(e.created_by) if e.created_by else None),
            "created_at": e.created_at.isoformat() if e.created_at else None,
        })
    return out


def _voided_export_nos(db: Session, period: SalaryPeriod, kind: str) -> list[int]:
    unlocks = _unlock_events(db, period.id)
    nos = []
    for e in _export_events(db, period.id):
        p = e.payload or {}
        if p.get("kind") == kind and _export_status(e, unlocks, period)[0] == "voided":
            nos.append(p.get("export_no"))
    return [n for n in nos if n is not None]


def _operator_name(db: Session, operator_id: Optional[int]) -> str:
    if not operator_id:
        return "未知"
    from app.auth.models import ArkUser

    name = db.query(ArkUser.real_name).filter(ArkUser.id == operator_id).scalar()
    return name or str(operator_id)


def _run_export(
    db: Session,
    period: SalaryPeriod,
    *,
    kind: str,
    payload: dict[str, Any],
    operator_id: Optional[int],
    expected: tuple[str, int],
    build,
) -> tuple[str, bytes]:
    """模块注释第 2 条与「并发」：事件与文件同一事务。

    expected = 取数时的 (status, status_version)，调用方在 load_export_rows 之前记下。
    build(export_no, now) -> (filename, bytes)。
    """
    now = beijing_now()
    try:
        # 拿批次行锁并发号。UPDATE 读的是最新已提交值，之后同事务的读能看到自己这次写
        db.execute(
            update(SalaryPeriod)
            .where(SalaryPeriod.id == period.id)
            .values(export_seq=SalaryPeriod.export_seq + 1)
            .execution_options(synchronize_session=False)
        )
        fresh = db.execute(
            select(SalaryPeriod.export_seq, SalaryPeriod.status, SalaryPeriod.status_version)
            .where(SalaryPeriod.id == period.id)
        ).one()
        if (fresh.status, fresh.status_version) != expected:
            raise SalaryExportError("批次状态刚刚被他人变更（锁定或解锁），请刷新页面后重新导出")
        export_no = fresh.export_seq
        event = SalaryPeriodEvent(
            period_id=period.id,
            event_type=EVENT_EXPORT,
            status_version=fresh.status_version,
            payload={**payload, "kind": kind},
            created_by=operator_id,
            created_at=now,
        )
        db.add(event)
        db.flush()
        filename, content = build(export_no, now)
        # 重新赋值整个 dict：JSON 列原地改字典 SQLAlchemy 感知不到
        event.payload = {
            **payload,
            "kind": kind,
            "export_no": export_no,
            "filename": filename,
            "sha256": hashlib.sha256(content).hexdigest(),
        }
        db.commit()
    except SalaryExportError:
        db.rollback()
        raise
    except Exception as exc:
        db.rollback()
        # 只记异常类型：openpyxl 等库的异常文本会带出单元格内容（姓名、备注）
        logger.warning("薪资导出失败 period=%s kind=%s: %s", period.id, kind, type(exc).__name__)
        print(f"[salary.export] failed period={period.id} kind={kind}: {type(exc).__name__}",
              flush=True)
        raise
    # 只记数量与金额：明文卡号、姓名都不进日志
    print(
        f"[salary.export] {period.year_month} kind={kind} no={export_no} "
        f"rows={payload.get('record_count')} plain_card={payload.get('plain_card')} by={operator_id}",
        flush=True,
    )
    return filename, content


# ---------------------------------------------------------------------------
# 工资表（双 sheet）
# ---------------------------------------------------------------------------

_FONT = "微软雅黑"
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 列宽取自 3 月人工表（xls 列宽单位 1/256 字符）
_SUMMARY_WIDTHS = [5.2, 7, 9, 10.8, 10.3, 23.8, 40, 11, 32]
_DETAIL_WIDTHS = [
    4.8, 10.8, 6, 10, 10.3, 12.2, 7, 8.3, 10.8, 7, 9, 9.5, 9.2, 10.7, 10.5, 8.3, 10.5,
    8.3, 10.5, 9.5, 10.5, 8.8, 14.5,
]

SUMMARY_HEADERS = ["序号", "部门", "姓名", "职务", "入职日期", "银行卡号", "请假时间", "实发工资", "备注"]
DETAIL_HEADERS = [
    "序号", "入职时间", "工号", "姓名", "部门", "职位", "应出天数", "实出天数", "应发工资",
    "奖励", "绩效", "工龄", "全勤奖", "小计", "社保", "公积金", "缺勤扣款", "其他款", "小计",
    "补贴项", "实发工资", "个税", "扣除个税实发",
]
# 明细表第 9 列（应发工资）起是金额列，合计行逐列求和
_DETAIL_MONEY_FROM = 8


def _detail_values(seq: int, row: ExportRow) -> list[Any]:
    r = row.record
    return [
        seq, row.hire_date, row.emp_no, row.name, row.dept_detail, row.position,
        r.due_days, r.actual_days, r.base_salary,
        r.bonus_final, r.performance_final, r.seniority_pay, r.attendance_bonus, r.add_subtotal,
        r.social_insurance, r.housing_fund, r.absence_deduction, r.other_final, r.deduct_subtotal,
        r.subsidy_final, r.net_salary, r.income_tax_amount, r.net_after_tax,
    ]


def _cell(ws, row: int, col: int, value: Any, *, size: int = 9, bold: bool = False,
          align: Alignment = _CENTER, fmt: Optional[str] = None, border: bool = True):
    if isinstance(value, str):
        # HR 从 Excel 粘进来的文本可能带  一类控制字符，openpyxl 会直接抛错
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=_FONT, size=size, bold=bold)
    c.alignment = align
    if border:
        c.border = _BORDER
    if fmt:
        c.number_format = fmt
    return c


def _sign_off(ws, row: int, cols: tuple[int, int, int]) -> None:
    for col, label in zip(cols, ("制表：", "审核：", "复核：")):
        _cell(ws, row, col, label, size=11, align=_LEFT, border=False)


def _stamp(ws, row: int, text: str, last_col: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = _cell(ws, row, 1, text, size=8, align=_LEFT, border=False)
    c.font = Font(name=_FONT, size=8, color="666666")


def _print_setup(ws, *, title_rows: str, header: str, footer: str) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows
    if header:
        ws.oddHeader.center.text = header
        ws.oddHeader.center.size = 14
        ws.oddHeader.center.color = "CC0000"
    ws.oddFooter.left.text = footer
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"
    ws.oddFooter.right.size = 8


def _write_summary(ws, title: str, groups, plain_card: bool, stamp: str, footer: str,
                   watermark: str) -> None:
    last = len(SUMMARY_HEADERS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    _cell(ws, 1, 1, title, size=14, bold=True, border=False)
    ws.row_dimensions[1].height = 46.5
    for i, h in enumerate(SUMMARY_HEADERS, start=1):
        _cell(ws, 2, i, h, size=11, bold=True)
    ws.row_dimensions[2].height = 30

    r = 3
    grand = Decimal("0")
    for group_name, items in groups:
        first = r
        subtotal = Decimal("0")
        for seq, row in enumerate(items, start=1):
            rec = row.record
            _cell(ws, r, 1, seq)
            _cell(ws, r, 2, group_name if seq == 1 else None)
            _cell(ws, r, 3, row.name)
            _cell(ws, r, 4, row.position)
            _cell(ws, r, 5, row.hire_date, fmt="yyyy/mm/dd")
            # 卡号写成文本：数字格式下 Excel 会把 19 位卡号变成科学计数、末几位清零
            _cell(ws, r, 6, _card_text(row, plain_card), fmt="@")
            _cell(ws, r, 7, rec.leave_remark, align=_LEFT)
            _cell(ws, r, 8, rec.net_salary, fmt="0")
            _cell(ws, r, 9, rec.remark_summary, align=_LEFT)
            ws.row_dimensions[r].height = 20
            subtotal += _dec(rec.net_salary)
            r += 1
        if r - 1 > first:
            ws.merge_cells(start_row=first, start_column=2, end_row=r - 1, end_column=2)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        _cell(ws, r, 1, "合计", size=10, bold=True)
        for col in range(2, 8):
            ws.cell(row=r, column=col).border = _BORDER
        _cell(ws, r, 8, subtotal, size=10, bold=True, fmt="0")
        _cell(ws, r, 9, None)
        grand += subtotal
        r += 1

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    _cell(ws, r, 1, "总计", size=10, bold=True)
    for col in range(2, 8):
        ws.cell(row=r, column=col).border = _BORDER
    _cell(ws, r, 8, grand, size=10, bold=True, fmt="0")
    _cell(ws, r, 9, None)
    r += 2
    _sign_off(ws, r, (2, 5, 8))
    _stamp(ws, r + 2, stamp, last)

    for i, w in enumerate(_SUMMARY_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    _print_setup(ws, title_rows="1:2", header=watermark, footer=footer)


def _write_detail(ws, rows: list[ExportRow], stamp: str, footer: str, watermark: str) -> None:
    last = len(DETAIL_HEADERS)
    for i, h in enumerate(DETAIL_HEADERS, start=1):
        _cell(ws, 1, i, h, size=10, bold=True)
    ws.row_dimensions[1].height = 33

    totals = [Decimal("0")] * last
    r = 2
    for seq, row in enumerate(rows, start=1):
        values = _detail_values(seq, row)
        for i, v in enumerate(values):
            if i == 1:
                fmt = "yyyy.m.d"
            elif i in (6, 7):
                fmt = "General"
            elif i >= _DETAIL_MONEY_FROM:
                fmt = "0.00"
            else:
                fmt = "@" if i == 2 else None
            _cell(ws, r, i + 1, v, fmt=fmt)
            if i >= _DETAIL_MONEY_FROM:
                totals[i] += _dec(v)
        ws.row_dimensions[r].height = 26
        r += 1

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=_DETAIL_MONEY_FROM)
    _cell(ws, r, 1, "合计", size=9, bold=True)
    for col in range(2, _DETAIL_MONEY_FROM + 1):
        ws.cell(row=r, column=col).border = _BORDER
    for i in range(_DETAIL_MONEY_FROM, last):
        _cell(ws, r, i + 1, totals[i], bold=True, fmt="0.00")
    r += 2
    # 制表/审核/复核的位置照 3 月人工表：D / I / R 列
    _sign_off(ws, r, (4, 9, 18))
    _stamp(ws, r + 2, stamp, last)

    for i, w in enumerate(_DETAIL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    _print_setup(ws, title_rows="1:1", header=watermark, footer=footer)


def _sheet_title(text: str) -> str:
    # Excel sheet 名上限 31 字符，且不能含 []:*?/\
    for ch in "[]:*?/\\":
        text = text.replace(ch, "_")
    return text[:31]


def build_workbook(
    period: SalaryPeriod,
    rows: list[ExportRow],
    groups: list[tuple[str, list[ExportRow]]],
    *,
    plain_card: bool,
    stamp: str,
    preview: bool,
) -> bytes:
    """生成与 3 月人工表同构的 xlsx。纯函数：不碰 session，便于单测。"""
    year, month = period_service.parse_year_month(period.year_month)
    watermark = "预览稿 · 批次未锁定 · 不得作为发放依据" if preview else ""
    summary_title = f"{year}年{month}月工资汇总表"
    detail_title = f"{year}年{month}月工资表（明细）"

    wb = Workbook()
    ws1 = wb.active
    ws1.title = _sheet_title(summary_title)
    shown_title = f"{summary_title}（预览稿·未锁定，不得作为发放依据）" if preview else summary_title
    # 页眉页脚有 255 字符上限，作废编号一多就超；完整脚注留在表底单元格里
    footer = stamp.split("｜", 1)[0]
    _write_summary(ws1, shown_title, groups, plain_card, stamp, footer, watermark)
    ws2 = wb.create_sheet(_sheet_title(detail_title))
    _write_detail(ws2, rows, stamp, footer, watermark)

    wb.properties.title = shown_title
    wb.properties.creator = "莱莎方舟·薪资模块"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt_dt(v: Optional[datetime]) -> str:
    return v.strftime("%Y-%m-%d %H:%M") if v else "-"


def _stamp_text(
    *, export_no: int, now: datetime, operator: str, period: SalaryPeriod,
    preview: bool, plain_card: bool, voided_nos: list[int], missing_cards: int = 0,
) -> str:
    # 第一段「方舟系统导出 第N次 时间」同时用作打印页脚，保持短
    parts = [f"方舟系统导出 第{export_no}次 {_fmt_dt(now)}", f"操作人 {operator}"]
    if preview:
        parts.append(f"批次状态：{period_service.STATUS_LABELS.get(period.status, period.status)}（未锁定，预览稿）")
    else:
        parts.append(f"批次锁定于 {_fmt_dt(period.confirmed_at)}")
    if voided_nos:
        nos = "、".join(str(n) for n in voided_nos)
        parts.append(f"本表替代第 {nos} 次导出（已因解锁作废）")
    if missing_cards:
        parts.append(f"{missing_cards} 人无银行卡号")
    parts.append("含完整银行卡号，请妥善保管" if plain_card else "银行卡号已脱敏")
    return "｜".join(parts)


def export_workbook(
    db: Session,
    period: SalaryPeriod,
    *,
    plain_card: bool,
    operator_id: Optional[int],
) -> tuple[str, bytes]:
    expected = (period.status, period.status_version)
    rows = load_export_rows(db, period)
    groups = group_rows(db, rows)
    missing_cards = len(_missing_card_names(rows))
    preview = not is_locked(period)
    operator = _operator_name(db, operator_id)
    voided_nos = [] if preview else _voided_export_nos(db, period, KIND_WORKBOOK)
    total = sum((_dec(r.record.net_salary) for r in rows), Decimal("0"))

    def build(export_no: int, now: datetime) -> tuple[str, bytes]:
        stamp = _stamp_text(
            export_no=export_no, now=now, operator=operator, period=period,
            preview=preview, plain_card=plain_card, voided_nos=voided_nos,
            missing_cards=missing_cards,
        )
        content = build_workbook(period, rows, groups, plain_card=plain_card,
                                 stamp=stamp, preview=preview)
        tag = "预览稿_" if preview else ""
        return f"{period.year_month}工资表_{tag}第{export_no}次导出.xlsx", content

    return _run_export(
        db, period, kind=KIND_WORKBOOK, operator_id=operator_id, expected=expected, build=build,
        payload={
            "locked": not preview,
            "plain_card": plain_card,
            "entity": None,
            "record_count": len(rows),
            "total_amount": str(total),
        },
    )


# ---------------------------------------------------------------------------
# 银行代发盘
# ---------------------------------------------------------------------------

BANK_HEADERS = ["姓名", "卡号", "金额"]


def _bank_amount(row: ExportRow) -> Optional[Decimal]:
    """代发金额 = 扣除个税实发（真正打到卡上的钱）。3 月个税全 0，与实发相同。"""
    v = row.record.net_after_tax
    return Decimal(v) if v is not None else None


def _entity_of(row: ExportRow) -> str:
    return (row.insurance_entity or "").strip() or NO_ENTITY


def _bank_check(rows: list[ExportRow]) -> tuple[list[ExportRow], list[dict], list[str]]:
    """(可发的行, 阻断问题, 实发为 0 不进盘的人)。

    金额 > 0 二次校验（§6 负数实发拦截的最后一道）：负数或没算出来的金额不能进盘，
    也不能悄悄跳过——跳过等于少发一个人且没人发现，所以是阻断。实发恰好为 0 的人
    没有钱要打，网银批量代发也不接受 0 元行，列出来告诉出纳即可。
    """
    payable, issues, zero = [], [], []
    for row in rows:
        amount = _bank_amount(row)
        who = row.name or f"员工#{row.employee_id}"
        if amount is None:
            issues.append({"name": who, "problem": "扣除个税实发为空（批次可能需要重新计算）"})
            continue
        if amount < 0:
            issues.append({"name": who, "problem": f"扣除个税实发为负数（{amount}）"})
            continue
        if amount == 0:
            zero.append(who)
            continue
        if not _plain_card(row):
            issues.append({"name": who, "problem": "锁定快照里没有银行卡号，或卡号无法解密"})
            continue
        payable.append(row)
    return payable, issues, zero


def bank_overview(db: Session, period: SalaryPeriod) -> dict[str, Any]:
    """银行盘预检：按参保主体分组的人数/金额/阻断问题。不解密出站，只给结论。"""
    if not is_locked(period):
        return {"available": False, "reason": "批次锁定后才能生成银行代发盘", "groups": []}
    try:
        rows = load_export_rows(db, period)
    except SalaryExportError as exc:
        return {"available": False, "reason": str(exc), "groups": []}
    by_entity: dict[str, list[ExportRow]] = {}
    for row in rows:
        by_entity.setdefault(_entity_of(row), []).append(row)
    groups = []
    for entity in sorted(by_entity, key=lambda e: (e == NO_ENTITY, e)):
        payable, issues, zero = _bank_check(by_entity[entity])
        groups.append({
            "entity": entity,
            "count": len(payable),
            "total_amount": str(sum((_bank_amount(r) for r in payable), Decimal("0"))),
            "issues": issues,
            "zero_names": zero,
        })
    return {"available": True, "reason": None, "groups": groups}


def build_bank_csv(rows: list[ExportRow]) -> bytes:
    """姓名 / 卡号 / 金额三列。UTF-8 带 BOM：Excel 直接打开中文不乱码。"""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(BANK_HEADERS)
    for row in rows:
        writer.writerow([row.name or "", _plain_card(row), f"{_bank_amount(row):.2f}"])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def export_bank(
    db: Session,
    period: SalaryPeriod,
    *,
    entity: Optional[str],
    operator_id: Optional[int],
) -> tuple[str, bytes]:
    """entity 为空 = 合并成一个盘；否则只出该参保主体的盘（默认拆两盘，§8.3 第 12 条）。"""
    if not is_locked(period):
        raise SalaryExportError("批次锁定后才能生成银行代发盘")
    expected = (period.status, period.status_version)
    rows = load_export_rows(db, period)
    target = (entity or "").strip() or None
    if target is not None:
        rows = [r for r in rows if _entity_of(r) == target]
        if not rows:
            raise SalaryExportError(f"本批次没有参保主体为「{target}」的人")
    payable, issues, _zero = _bank_check(rows)
    if issues:
        detail = "；".join(f"{i['name']}：{i['problem']}" for i in issues[:10])
        more = f" 等 {len(issues)} 人" if len(issues) > 10 else ""
        raise SalaryExportError(f"以下人员不能进代发盘，请处理后再导出——{detail}{more}")
    if not payable:
        raise SalaryExportError("没有需要代发的人（实发均为 0）")
    total = sum((_bank_amount(r) for r in payable), Decimal("0"))
    label = target or "全部"

    def build(export_no: int, now: datetime) -> tuple[str, bytes]:
        return (f"{period.year_month}银行代发_{label}_第{export_no}次导出.csv",
                build_bank_csv(payable))

    return _run_export(
        db, period, kind=KIND_BANK, operator_id=operator_id, expected=expected, build=build,
        payload={
            "locked": True,
            "plain_card": True,
            "entity": target,
            "record_count": len(payable),
            "total_amount": str(total),
        },
    )


# ---------------------------------------------------------------------------
# 概览（导出面板一次拿齐）
# ---------------------------------------------------------------------------

def export_overview(db: Session, period: SalaryPeriod) -> dict[str, Any]:
    """可导项用与导出相同的取数函数判定：预检说能导、点了却 400，比不预检还糟。"""
    locked = is_locked(period)
    workbook: dict[str, Any] = {"available": False, "preview": not locked, "reason": None,
                                "missing_card_names": []}
    try:
        rows = load_export_rows(db, period)
    except SalaryExportError as exc:
        workbook["reason"] = str(exc)
    else:
        workbook["available"] = True
        workbook["missing_card_names"] = _missing_card_names(rows)
    history = export_history(db, period)
    return {
        "locked": locked,
        "workbook": workbook,
        "bank": bank_overview(db, period),
        "history": history,
        "voided_count": sum(1 for h in history if h["status"] == "voided"),
    }


def content_disposition(filename: str) -> str:
    """RFC 5987 文件名：中文名浏览器原样保存，前端 utils/download 认 filename*。"""
    return f"attachment; filename*=UTF-8''{quote(filename, safe='')}"
